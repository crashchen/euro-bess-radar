"""Generate real Step 3C exports, read saved cells, and render PDFs with Poppler.

Run from the repository root:
PYTHONPATH=. .venv/bin/python outputs/step3c-2026-09-16/export-check/export_probe.py
"""
from __future__ import annotations

import json
import math
from pathlib import Path
import subprocess

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.analytics import calculate_average_price, compare_zones
from src.export import export_comparison_to_bytes, export_to_bytes, export_to_pdf_bytes
from tests.test_step3a_display_contract import _export_args


OUTPUT = Path("outputs/step3c-2026-09-16/export-check")
OUTPUT.mkdir(parents=True, exist_ok=True)

before = pd.date_range("2025-09-01", "2025-10-01", inclusive="left", freq="h", tz="Europe/Berlin")
after = pd.date_range("2025-10-01", "2025-10-11", inclusive="left", freq="15min", tz="Europe/Berlin")
mixed = pd.DataFrame(
    {"price_eur_mwh": np.r_[np.full(len(before), 10.0), np.full(len(after), 100.0)]},
    index=before.append(after).tz_convert("UTC").rename("timestamp"),
)
partial = mixed.copy(deep=True)
partial.iloc[5, 0] = np.nan
partial.iloc[len(before) + 10, 0] = np.inf
partial.iloc[len(before) + 11, 0] = -np.inf
gap = pd.DataFrame(
    {"price_eur_mwh": np.full(47, 50.0)},
    index=pd.date_range("2026-01-01", periods=48, freq="h", tz="UTC").delete(30).rename("timestamp"),
)
unknown = pd.DataFrame(
    {"price_eur_mwh": [10.0, 100.0, 10.0]},
    index=pd.date_range("2026-01-01", periods=3, freq="2h", tz="UTC").rename("timestamp"),
)
samples = {
    "mixed": (mixed, 32.50, 960.0, None),
    "partial-finite": (partial, 31140 / 958.5, 958.5, None),
    "gap": (gap, None, None, "gap or a cadence change"),
    "unknown-cadence": (unknown, None, None, "not a supported native"),
}


def saved_cell(cell):
    return {
        "coordinate": cell.coordinate, "value": cell.value,
        "data_type": cell.data_type, "number_format": cell.number_format,
        "wrap_text": cell.alignment.wrap_text,
    }


results = {}
for name, (frame, expected, covered, reason_needle) in samples.items():
    original = frame.copy(deep=True)
    args = _export_args(frame)
    xlsx_path = OUTPUT / f"{name}.xlsx"
    pdf_path = OUTPUT / f"{name}.pdf"
    xlsx_path.write_bytes(export_to_bytes(**args))
    pdf_path.write_bytes(export_to_pdf_bytes(**args))
    pd.testing.assert_frame_equal(frame, original)
    ws = load_workbook(xlsx_path)["Summary"]
    pairs = {row[0].value: row[1] for row in ws.iter_rows(min_col=1, max_col=2) if row[0].value}
    average = pairs["Avg Price (EUR/MWh)"]
    if expected is not None:
        assert average.data_type == "n"
        assert math.isclose(average.value, round(expected, 2), rel_tol=0, abs_tol=1e-12)
        assert average.number_format == "#,##0.00"
        basis = pairs["Avg Price Basis"]
        assert f"{covered:,.2f} of 960.00 delivery hours" in basis.value
        if name == "partial-finite":
            assert "99.8%" in basis.value
            assert "excluded, not counted as zero" in basis.value
        assert "Avg Price Unavailable Because" not in pairs
    else:
        assert average.value == "n/a" and average.data_type == "s"
        basis = pairs["Avg Price Unavailable Because"]
        assert reason_needle in basis.value
        assert "Avg Price Basis" not in pairs
    text_path = OUTPUT / f"{name}.pdf.txt"
    subprocess.run(["pdftotext", "-layout", str(pdf_path), str(text_path)], check=True)
    text = text_path.read_text()
    normalized_text = " ".join(text.split())
    expected_token = f"{expected:.2f}" if expected is not None else "n/a"
    assert f"Avg Price (EUR/MWh) {expected_token}" in normalized_text
    assert " ".join(str(basis.value).split()) in normalized_text
    subprocess.run(["pdftoppm", "-scale-to", "1800", "-png", str(pdf_path), str(OUTPUT / name)], check=True)
    stats = calculate_average_price(frame)
    results[name] = {
        "expected_average": expected, "expected_finite_hours": covered,
        "average_cell": saved_cell(average), "basis_or_reason_cell": saved_cell(basis),
        "calculation": {key: None if isinstance(value, float) and not math.isfinite(value) else value for key, value in stats.items()},
        "pdf": str(pdf_path), "xlsx": str(xlsx_path), "pdf_text": str(text_path),
        "rendered_pages": [str(path) for path in sorted(OUTPUT.glob(f"{name}-[0-9]*.png"))],
    }

comparison = compare_zones({"DE_LU": mixed, "FR": partial, "IT_NORD": gap})
comparison_path = OUTPUT / "comparison.xlsx"
comparison_path.write_bytes(export_comparison_to_bytes(comparison))
ws = load_workbook(comparison_path)["Zone Comparison"]
columns = {cell.value: cell.column for cell in ws[1]}
comparison_rows = {}
for cells in ws.iter_rows(min_row=2):
    zone = cells[columns["Zone"] - 1].value
    avg = cells[columns["Avg Price (EUR/MWh)"] - 1]
    coverage = cells[columns["Avg Price Coverage %"] - 1]
    reason = cells[columns["Avg Price Unavailable Because"] - 1]
    comparison_rows[zone] = {"average": saved_cell(avg), "coverage": saved_cell(coverage), "reason": saved_cell(reason)}
    if zone == "IT_NORD":
        assert avg.value == "n/a" and coverage.value == "n/a"
        assert "gap or a cadence change" in reason.value
    else:
        assert avg.data_type == coverage.data_type == "n"
        assert math.isclose(avg.value, 32.50 if zone == "DE_LU" else 32.49, abs_tol=1e-12)
        assert math.isclose(coverage.value, 1.0 if zone == "DE_LU" else 0.9984, abs_tol=1e-12)
        assert coverage.number_format == "0.0%"
        assert reason.value is None
results["comparison"] = {"path": str(comparison_path), "rows": comparison_rows}
(OUTPUT / "results.json").write_text(json.dumps(results, indent=2, allow_nan=False) + "\n")
print(json.dumps(results, indent=2, allow_nan=False))
