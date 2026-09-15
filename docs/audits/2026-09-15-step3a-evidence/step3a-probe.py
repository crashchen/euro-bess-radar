"""Step 3A probe: unavailable durations and defective price indices, end to end.

Run from the repository root against any revision. It prints what each surface
shows, so the pre-change and post-change behaviour can be compared directly:

    PYTHONPATH=. .venv/bin/python docs/audits/2026-09-15-step3a-evidence/step3a-probe.py

Synthetic data only. It writes nothing to the price cache; the Excel and PDF
artefacts are written to a temporary directory and removed on exit.
"""

from __future__ import annotations

import contextlib
import re
import tempfile
import zlib
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.analytics import (
    calculate_daily_spreads,
    calculate_negative_price_hours,
    calculate_spread_percentiles,
    estimate_annual_arbitrage_revenue,
)
from src.export import export_to_bytes, export_to_pdf_bytes


def unverifiable_frame() -> pd.DataFrame:
    """A complete negative hourly day, then a day whose grid cannot be inferred."""
    day_one = pd.date_range("2026-01-01", periods=24, freq="h", tz="UTC")
    day_two = pd.DatetimeIndex(
        pd.to_datetime(["2026-01-02T00:00Z", "2026-01-02T01:00Z", "2026-01-02T03:00Z"]),
    )
    index = day_one.append(day_two).rename("timestamp")
    return pd.DataFrame(
        {"price_eur_mwh": np.r_[np.full(24, -50.0), np.full(3, 60.0)]}, index=index,
    )


def pdf_text(data: bytes) -> str:
    chunks: list[str] = []
    for stream in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", data, re.S):
        raw = stream.group(1)
        with contextlib.suppress(zlib.error):
            raw = zlib.decompress(raw)
        chunks += [m.group(0)[1:-1].decode("latin-1") for m in
                   re.finditer(rb"\((?:\\.|[^\\()])*\)", raw)]
    return re.sub(r"\\([()\\])", r"\1", " ".join(chunks))


def export_args(frame: pd.DataFrame) -> dict:
    daily = calculate_daily_spreads(frame)
    return {
        "zone": "DE_LU", "price_df": frame, "daily_spreads": daily,
        "monthly_spreads": pd.DataFrame(),
        "percentiles": calculate_spread_percentiles(daily),
        "revenue_estimate": estimate_annual_arbitrage_revenue(daily),
        "negative_stats": calculate_negative_price_hours(frame), "tz": "UTC",
    }


def probe_negative_hours(out_dir: Path) -> None:
    frame = unverifiable_frame()
    stats = calculate_negative_price_hours(frame)
    print("[1] unverifiable grid, 24 negative hourly intervals + one unknown day")
    print(f"    negative_hours          = {stats['negative_hours']!r}")
    print(f"    negative_intervals      = {stats['negative_intervals']!r}")
    print(f"    pct_negative            = {stats['pct_negative']!r}")
    print(f"    negative_hours_reason   = {stats.get('negative_hours_reason')!r}")

    args = export_args(frame)
    xlsx = out_dir / "unverifiable.xlsx"
    xlsx.write_bytes(export_to_bytes(**args))
    ws = load_workbook(BytesIO(xlsx.read_bytes()))["Summary"]
    pairs = {r[0].value: r[1].value for r in ws.iter_rows(min_col=1, max_col=2) if r[0].value}
    print("    Excel 'Negative Price Hours'                  =", repr(pairs.get("Negative Price Hours")))
    print("    Excel 'Negative Price Hours Unavailable Because' =",
          repr(pairs.get("Negative Price Hours Unavailable Because")))
    print("    Excel 'Negative Price Intervals'              =", repr(pairs.get("Negative Price Intervals")))

    pdf = out_dir / "unverifiable.pdf"
    pdf.write_bytes(export_to_pdf_bytes(**args))
    text = pdf_text(pdf.read_bytes())
    start = text.find("Negative Price Hours")
    print("    PDF summary excerpt =", repr(text[start:start + 170]))
    print("    PDF contains literal 'nan' =", "nan" in text.lower())

    control = export_args(pd.DataFrame(
        {"price_eur_mwh": np.full(48, -50.0)},
        index=pd.date_range("2026-01-01", periods=48, freq="h", tz="UTC").rename("timestamp"),
    ))
    ws = load_workbook(BytesIO(export_to_bytes(**control)))["Summary"]
    pairs = {r[0].value: r[1].value for r in ws.iter_rows(min_col=1, max_col=2) if r[0].value}
    print("    control (48 verifiable negative hours): Excel cell =",
          repr(pairs.get("Negative Price Hours")),
          "| reason row present =", "Negative Price Hours Unavailable Because" in pairs)


def probe_market_page() -> None:
    from streamlit.testing.v1 import AppTest

    print("\n[2] market page with a defective price index")
    for defect in ("duplicate", "unsorted", "nat", "clean"):
        app = AppTest.from_function(_page_app, args=(defect,)).run(timeout=60)
        exception = app.exception[0].message if app.exception else None
        try:
            traces = app.session_state["trace_names"]
        except (KeyError, AttributeError):
            traces = "<not reached>"
        print(f"    {defect:<9} exception={exception!r}")
        print(f"              price_ts traces={traces!r}")
        print(f"              errors={[e.value[:70] for e in app.error]}")
        print(f"              warnings={[w.value[:70] for w in app.warning]}")


def _page_app(defect: str) -> None:
    # AppTest re-executes this body in a fresh module, so it imports its own
    # dependencies and builds its own frame rather than closing over globals.
    import numpy as np
    import pandas as pd
    import streamlit as st

    from src.analytics import (
        calculate_daily_spreads,
        calculate_negative_price_hours,
        calculate_spread_percentiles,
    )
    from src.pages.market_overview import render

    index = pd.date_range("2026-01-01", periods=48, freq="h", tz="UTC")
    values = np.arange(48.0)
    if defect == "duplicate":
        order = [*range(48), 10]
        index, values = index[order].sort_values(), np.sort(values[order])
    elif defect == "unsorted":
        order = [*range(24, 48), *range(24)]
        index, values = index[order], values[order]
    elif defect == "nat":
        index = pd.DatetimeIndex([*index[:47], pd.NaT])
    frame = pd.DataFrame({"price_eur_mwh": values}, index=index.rename("timestamp"))
    daily = calculate_daily_spreads(frame)
    figures: dict[str, object] = {"price_ts": "stale-figure-from-an-earlier-state"}
    render("DE_LU", frame, daily, calculate_spread_percentiles(daily),
           calculate_negative_price_hours(frame), 1, "UTC", "plotly_dark", figures)
    price_ts = figures.get("price_ts")
    st.session_state["trace_names"] = (
        [t.name for t in price_ts.data] if hasattr(price_ts, "data") else price_ts
    )


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as tmp:
        probe_negative_hours(Path(tmp))
    probe_market_page()
