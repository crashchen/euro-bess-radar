"""Generate production Step 3D export fixtures without accessing price caches.

PC and screening use real public solvers. Only the Cockpit comparison fixture
uses the explicitly named solver stubs from the dedicated display test.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from src.export import (
    cockpit_tables_to_excel,
    export_to_bytes,
    export_to_pdf_bytes,
    project_case_to_excel,
)
from tests.test_step3a_display_contract import _export_args
from tests.test_step3d_cockpit_disclosure import _compute_bundle
from tests.test_step3d_project_case_disclosure import synthetic_reserve_result


def _screening_revenue_app(zone, timezone):
    import datetime as dt
    from unittest.mock import patch

    import pandas as pd
    import streamlit as st

    from src.ancillary import _build_standard_frame
    from src.analytics import estimate_annual_arbitrage_revenue
    from src.pages import revenue_estimation as page
    from tests.test_step3a_display_contract import _export_args

    day = dt.date(2026, 3, 29)
    lo = pd.Timestamp(day, tz=timezone)
    hi = lo + pd.DateOffset(days=1)
    index = pd.date_range(lo, hi, freq="15min", inclusive="left").tz_convert("UTC")
    frame = pd.DataFrame({"price_eur_mwh": 40.0}, index=index)
    args = _export_args(frame)
    args["revenue_estimate"] = estimate_annual_arbitrage_revenue(
        args["daily_spreads"], power_mw=1.0, duration_hours=1,
        roundtrip_efficiency=1.0, capture_rate=1.0,
    )
    st.session_state["ancillary_df"] = (
        pd.concat([
            _build_standard_frame(index, "FCR", zone, capacity=20.0),
            _build_standard_frame(index, "aFRR Up", zone, capacity=10.0),
            _build_standard_frame(index, "mFRR Up", zone, energy=1.0),
        ]) if zone == "DE_LU" else _build_standard_frame(
            index, "FCR-N", zone, capacity=20.0,
        )
    )
    st.session_state["ancillary_zone"] = zone
    st.session_state["ancillary_dates"] = (str(day), str(day))
    export = args["revenue_estimate"].copy()
    # Unrelated subpanels are skipped. Actual capacity aggregation and the
    # production joint solver run on the synthetic zone's physical day.
    with patch.object(page, "_render_intraday_uplift_section"), \
         patch.object(page, "_render_sensitivity_table"), \
         patch.object(page, "_render_revenue_risk_analysis"), \
         patch.object(page, "render_project_case_panel"):
        page.render(
            primary_zone=zone, primary_df=frame,
            daily_spreads=args["daily_spreads"], monthly_spreads=args["monthly_spreads"],
            percentiles=args["percentiles"], revenue=args["revenue_estimate"],
            start_date=day, end_date=day, power_mw=1.0, duration_hours=1,
            efficiency=1.0, capture_rate=1.0, capex_eur_kwh=0.0, use_lp_dispatch=False,
            zone_tz=timezone, chart_template="plotly_dark", report_figures={},
            export_revenue=export,
        )
    st.session_state["test_export"] = export
    st.session_state["test_frame"] = frame


def daily_frame(timezone):
    lo = pd.Timestamp("2026-03-29", tz=timezone)
    index = pd.date_range(lo, lo + pd.DateOffset(days=1), freq="15min", inclusive="left")
    return pd.DataFrame(
        {"price_eur_mwh": 40.0}, index=index.tz_convert("UTC").rename("timestamp")
    )


def write_book(out, filename, data, views, report):
    path = out / filename
    path.write_bytes(data)
    book = load_workbook(path)
    entries = []
    for sheet in book:
        cells = []
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                cells.append({
                    "cell": cell.coordinate, "value": cell.value,
                    "type": cell.data_type, "number_format": cell.number_format,
                    "wrap": cell.alignment.wrap_text,
                    "row_height": sheet.row_dimensions[cell.row].height,
                })
        entries.append({
            "sheet": sheet.title, "cells": cells,
            "column_widths": {k: v.width for k, v in sheet.column_dimensions.items()},
        })
    (out / f"{path.stem}-cells.json").write_text(json.dumps(entries, indent=2, default=str) + "\n")
    for sheetname, label, column in views:
        sheet = book[sheetname]
        if label:
            row = next(c.row for c in sheet["A"] if c.value == label)
            start, end = max(1, row - 2), min(sheet.max_row, row + 3)
        else:
            start, end = 1, sheet.max_row
        report["renders"].append({
            "file": filename, "sheet": sheetname,
            "range": f"A{start}:{column}{end}",
            "png": f"{path.stem}-{sheetname.lower().replace(' ', '-')}.png",
        })
    report["artifacts"][filename] = {
        "sha256": hashlib.sha256(data).hexdigest(),
        "sheets": book.sheetnames,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--out", type=Path,
        default=Path(tempfile.gettempdir()) / "euro-bess-step3d-exports",
    )
    out = parser.parse_args().out
    out.mkdir(parents=True, exist_ok=True)
    report = {"artifacts": {}, "renders": [], "known_answers": {}, "limits": [
        "Excel images are ArtifactTool read-only import renders, not Microsoft Excel screenshots.",
        "PDF images are Poppler renders of production exporter bytes; no charts requested in these focused summary fixtures.",
        "Cockpit fixture uses the explicit solver stubs in test_step3d_cockpit_disclosure._compute_bundle; it proves display/export preservation, not economics.",
    ]}
    result = synthetic_reserve_result()
    report["known_answers"]["project_case"] = {
        "real_solver": True, "zone": "DE_LU", "product": "FCR [symmetric]",
        "physical_hours": 23, "nominal_hours": 24, "availability": 0.95,
        "daily_realised_cash_eur": result.provenance["strategy_run_result"]["daily_realised_cash_series"][0][1],
        "input_fingerprint": result.input_fingerprint,
    }
    pcviews = [("Project Case NPVs", "Reserve Capacity Settlement Basis", "C")]
    write_book(out, "project-case-dst.xlsx", project_case_to_excel(result), pcviews, report)
    appended = _export_args(daily_frame("Europe/Helsinki"))
    appended.update(zone="FI", tz="Europe/Helsinki")
    write_book(out, "fi-report-with-de-project-case.xlsx",
               export_to_bytes(**appended, project_case_result=result), pcviews, report)
    for zone, timezone, expected in [
        ("DE_LU", "Europe/Berlin", 30 * 0.95 * 23),
        ("FI", "Europe/Helsinki", 20 * 0.95 * 23),
    ]:
        app = AppTest.from_function(_screening_revenue_app, args=(zone, timezone)).run(timeout=60)
        assert not app.exception, [e.message for e in app.exception]
        revenue = app.session_state["test_export"]
        assert revenue["joint_cooptimized_capacity_eur"] == pytest.approx(expected * 365.25)
        basis = revenue["joint_capacity_settlement_basis"]
        if zone == "FI":
            assert "nominal 4h" not in basis
        args = _export_args(daily_frame(timezone))
        args.update(zone=zone, tz=timezone, revenue_estimate=revenue)
        stem = f"{zone.lower()}-screening-spring"
        write_book(out, f"{stem}.xlsx", export_to_bytes(**args),
                   [("Summary", "Joint MILP Capacity Settlement Basis", "B")], report)
        pdf = export_to_pdf_bytes(**args)
        path = out / f"{stem}.pdf"
        path.write_bytes(pdf)
        subprocess.run(["pdftoppm", "-png", "-r", "120", str(path), str(out / stem)], check=True)
        subprocess.run(["pdftotext", "-layout", str(path), str(out / f"{stem}-pdf.txt")], check=True)
        report["artifacts"][path.name] = {"sha256": hashlib.sha256(pdf).hexdigest()}
        report["known_answers"][zone] = {
            "real_solver": True, "physical_hours": 23, "availability": 0.95,
            "daily_capacity_eur": expected,
            "annual_capacity_eur": revenue["joint_cooptimized_capacity_eur"],
            "joint_da_eur": revenue["joint_cooptimized_da_eur"], "basis": basis,
        }
    with pytest.MonkeyPatch.context() as monkeypatch:
        bundle = _compute_bundle(monkeypatch, assumptions=None)
    derived = bundle["derived"]
    write_book(out, "cockpit-forecast-stub-no-input-assumptions.xlsx",
               cockpit_tables_to_excel(derived["export_tables"], assumptions=derived["export_assumptions"]),
               [("Strategy comparison", None, "F"), ("Assumptions", None, "E")], report)
    report["known_answers"]["cockpit_stub"] = {
        "real_solver": False, "caller_assumptions": None,
        "reserve_window_eur": 437.0, "policy_value_delta_eur": -3.0,
        "basis": derived["capacity_settlement"],
    }
    (out / "export-verification.json").write_text(json.dumps(report, indent=2, default=str) + "\n")
    print(json.dumps({"artifacts": list(report["artifacts"]), "known_answers": report["known_answers"]}, indent=2, default=str))


if __name__ == "__main__":
    main()
