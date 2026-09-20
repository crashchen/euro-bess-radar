"""Regenerate production exports and verify the unchanged row-based statistics.

Run from the repository root with PYTHONPATH=. .venv/bin/python <this file>.
"""
from __future__ import annotations

import json
from pathlib import Path
import subprocess

from openpyxl import load_workbook

from src.analytics import compare_zones
from src.export import export_comparison_to_bytes, export_to_bytes, export_to_pdf_bytes
from tests.test_step3a_display_contract import _export_args
from tests.test_step3c_average_price import _cutover_sample, _gap_sample


folder = Path(__file__).resolve().parent
mixed = _cutover_sample()
args = _export_args(mixed)
(folder / "mixed.xlsx").write_bytes(export_to_bytes(**args))
(folder / "mixed.pdf").write_bytes(export_to_pdf_bytes(**args))
comparison = compare_zones({"DE_LU": mixed, "FR": _gap_sample()})
(folder / "comparison.xlsx").write_bytes(export_comparison_to_bytes(comparison))

summary = load_workbook(folder / "mixed.xlsx")["Summary"]
cells = {row[0].value: (row[0], row[1]) for row in summary.iter_rows(min_col=1, max_col=2) if row[0].value}
label, median = cells["Median Price (row-based, EUR/MWh)"]
assert cells["Avg Price (EUR/MWh)"][1].value == 32.50
assert median.value == float(mixed["price_eur_mwh"].median()) == 100.0
assert median.data_type == "n" and median.number_format == "#,##0.00"
assert label.alignment.wrap_text is True
assert summary.row_dimensions[label.row].height == 30.0

sheet = load_workbook(folder / "comparison.xlsx")["Zone Comparison"]
header = {cell.value: cell.column for cell in sheet[1]}
std = sheet.cell(row=2, column=header["Std Dev (row-based)"])
assert std.data_type == "n" and std.number_format == "#,##0.00"
assert std.value == round(float(mixed["price_eur_mwh"].std()), 2) == 44.55

subprocess.run(["pdftotext", "-layout", str(folder / "mixed.pdf"), str(folder / "mixed.pdf.txt")], check=True)
text = " ".join((folder / "mixed.pdf.txt").read_text().split())
assert "Median Price (row-based, EUR/MWh) 100.00" in text
subprocess.run(["pdftoppm", "-scale-to", "1800", "-png", str(folder / "mixed.pdf"), str(folder / "mixed")], check=True)
report = {
    "avg_price": cells["Avg Price (EUR/MWh)"][1].value,
    "median": {"label": label.value, "value": median.value, "data_type": median.data_type, "format": median.number_format, "wrap_text": label.alignment.wrap_text, "row_height": summary.row_dimensions[label.row].height},
    "std": {"label": "Std Dev (row-based)", "value": std.value, "data_type": std.data_type, "format": std.number_format},
    "pdf_text_verified": True,
}
(folder / "disclosure-inspection.json").write_text(json.dumps(report, indent=2) + "\n")
print(json.dumps(report, indent=2))
