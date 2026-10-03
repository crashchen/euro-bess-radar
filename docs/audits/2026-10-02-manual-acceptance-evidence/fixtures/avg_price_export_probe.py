"""Generate and inspect the production exports for every avg-price fixture.

Execs the fixture half of avg_price_harness.py (everything above the
module-level `scratch = ...` launcher) so the frames are byte-for-byte the
ones the browser harness renders. Cache paths are patched to a scratch dir.
"""

from __future__ import annotations

import io
import subprocess
import sys
from pathlib import Path
from tempfile import mkdtemp
from unittest.mock import patch

import openpyxl

HERE = Path(__file__).parent
OUT = HERE.parent / "downloads" / "synthetic"
src = (HERE / "avg_price_harness.py").read_text()
ns: dict = {"__name__": "fixture_probe"}
exec(compile(src.split("\nscratch = Path(")[0], "avg_price_harness.py", "exec"), ns)

from src import config, data_ingestion, export  # noqa: E402
from src.analytics import compare_zones  # noqa: E402
from src.config import ZONE_TIMEZONES  # noqa: E402
from src.export import export_comparison_to_bytes  # noqa: E402

scratch = Path(mkdtemp(prefix="avg-price-probe-"))
KEYS = ("Avg Price", "Median Price", "Negative Price Hours", "Negative Price Interval")
with (
    patch.object(config, "CACHE_DIR", scratch),
    patch.object(config, "DB_PATH", scratch / "synthetic.db"),
    patch.object(data_ingestion, "CACHE_DIR", scratch),
    patch.object(data_ingestion, "DB_PATH", scratch / "synthetic.db"),
    patch.object(export, "CACHE_DIR", scratch),
):
    for k, (name, df) in enumerate(ns["fixtures"]().items()):
        tag = name.split()[0]
        print(f"\n===== {name}  rows={len(df)}")
        other = ns["_frame"](ns["_local"]("2025-09-01", "2025-09-06", "h"), 60.0)
        comp = compare_zones(
            {"DE_LU": df, "FR": other}, zone_timezones=ZONE_TIMEZONES, duration_hours=1,
            capture_rate=0.7, roundtrip_efficiency=0.88, power_mw=10,
            use_lp_dispatch=False, capex_eur_kwh=0,
        )
        cx = export_comparison_to_bytes(comp)
        (OUT / f"f{tag}_comparison.xlsx").write_bytes(cx)
        ws = openpyxl.load_workbook(io.BytesIO(cx)).active
        hdr = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "DE_LU":
                d = {h: (c.value, c.data_type) for h, c in zip(hdr, row)}
                print("  comparison:", {h: d[h] for h in hdr if "Avg Price" in str(h) or h == "Std Dev (row-based)"})
        from src.analytics import (
            calculate_daily_spreads, calculate_monthly_spreads_from_daily,
            calculate_negative_price_hours, calculate_spread_percentiles,
            estimate_annual_arbitrage_revenue, filter_to_complete_local_days,
        )
        try:
            daily = calculate_daily_spreads(df, tz=ns["TZ"], duration_hours=1)
            monthly = calculate_monthly_spreads_from_daily(daily)
            pct = calculate_spread_percentiles(daily)
            neg = calculate_negative_price_hours(filter_to_complete_local_days(df, tz=ns["TZ"]))
            rev = estimate_annual_arbitrage_revenue(daily, power_mw=10, duration_hours=1,
                                                    roundtrip_efficiency=0.88, capture_rate=0.7)
            rev.update(power_mw=10, duration_hours=1, roundtrip_efficiency=0.88,
                       capture_basis="sidebar capture haircut applied to screening revenue")
            kw = dict(zone="DE_LU", price_df=df, daily_spreads=daily, monthly_spreads=monthly,
                      percentiles=pct, revenue_estimate=rev, negative_stats=neg, tz=ns["TZ"])
            xb = export.export_to_bytes(**kw, project_case_result=None)
            pb = export.export_to_pdf_bytes(**kw, figures=None)
        except Exception as exc:  # report, do not hide
            print("  EXPORT ERROR:", type(exc).__name__, exc)
            continue
        (OUT / f"f{tag}_report.xlsx").write_bytes(xb)
        (OUT / f"f{tag}_report.pdf").write_bytes(pb)
        ws = openpyxl.load_workbook(io.BytesIO(xb))["Summary"]
        for r in ws.iter_rows():
            if r[0].value and any(s in str(r[0].value) for s in KEYS):
                print(f"  xlsx: {r[0].value} = {r[1].value!r} [{r[1].data_type}]")
        txt = subprocess.run(["pdftotext", "-layout", str(OUT / f"f{tag}_report.pdf"), "-"],
                             capture_output=True, text=True).stdout
        for line in txt.splitlines():
            if any(s in line for s in ("Avg Price", "Median Price", "Negative Price Hours")) or "nan" in line.lower():
                print("  pdf:", line.strip()[:200])
sys.exit(0)
