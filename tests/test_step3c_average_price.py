"""Step 3C: the window-wide Avg Price is weighted by verified delivery duration.

One shared calculation feeds the market page, the zone comparison and the
Excel/PDF summaries. A quarter-hour price counts for a quarter hour, so a
window that crosses the SDAC hourly-to-quarter-hour cutover no longer lets the
denser half dominate. Where the delivery duration cannot be verified the
average is ``n/a`` with a reason on every surface, never a guessed number.

Symbols introduced by this step are imported inside the cases that use them,
so the module still collects and fails on behaviour against the baseline.
"""

from __future__ import annotations

import json
from io import BytesIO

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from src.analytics import compare_zones
from src.export import export_comparison_to_bytes, export_to_bytes, export_to_pdf_bytes
from tests.test_step3a_display_contract import _export_args, _pdf_chunks, _pdf_text

_GAP_REASON = "gap or a cadence change"


# ── Samples ──────────────────────────────────────────────────────────────────

def _frame(index: pd.DatetimeIndex, values) -> pd.DataFrame:
    return pd.DataFrame(
        {"price_eur_mwh": np.asarray(values, dtype=float)},
        index=pd.DatetimeIndex(index).rename("timestamp"),
    )


def _cutover_sample() -> pd.DataFrame:
    """30 hourly days at 10, then 10 quarter-hour days at 100 (SDAC cutover)."""
    before = pd.date_range(
        "2025-09-01", "2025-10-01", inclusive="left", freq="h", tz="Europe/Berlin",
    )
    after = pd.date_range(
        "2025-10-01", "2025-10-11", inclusive="left", freq="15min", tz="Europe/Berlin",
    )
    return _frame(
        before.append(after).tz_convert("UTC"),
        np.r_[np.full(len(before), 10.0), np.full(len(after), 100.0)],
    )


def _uniform_sample(freq: str) -> pd.DataFrame:
    """40 days on one grid: the first 30 at 10 and the last 10 at 100."""
    per_day = {"h": 24, "30min": 48, "15min": 96}[freq]
    count = 40 * per_day
    index = pd.date_range("2026-01-01", periods=count, freq=freq, tz="UTC")
    return _frame(index, np.r_[np.full(count * 3 // 4, 10.0), np.full(count // 4, 100.0)])


def _civil_cutover_sample(tz: str) -> pd.DataFrame:
    """Two civil days before and one after the shared SDAC cutover instant."""
    cutover = pd.Timestamp("2025-09-30T22:00:00Z")
    before = pd.date_range(
        pd.Timestamp("2025-09-29", tz=tz).tz_convert("UTC"), cutover,
        inclusive="left", freq="h",
    )
    after = pd.date_range(
        cutover, pd.Timestamp("2025-10-02", tz=tz).tz_convert("UTC"),
        inclusive="left", freq="15min",
    )
    return _frame(
        before.append(after), np.r_[np.full(len(before), 10.0), np.full(len(after), 100.0)],
    )


def _gap_sample() -> pd.DataFrame:
    """Hourly prices with one delivery hour absent from the index entirely."""
    index = pd.date_range("2026-01-01", periods=48, freq="h", tz="UTC").delete(30)
    return _frame(index, np.full(47, 50.0))


def _defective_sample(defect: str) -> pd.DataFrame:
    from tests.test_step3a_display_contract import _defective

    return _defective(defect)


# ── Calculation contract ─────────────────────────────────────────────────────

def test_cutover_window_is_weighted_by_delivery_hours_not_rows() -> None:
    from src.analytics import calculate_average_price

    frame = _cutover_sample()
    stats = calculate_average_price(frame)
    assert stats["avg_price_eur_mwh"] == pytest.approx(32.50)
    assert stats["covered_hours"] == pytest.approx(30 * 24 + 10 * 24)
    assert stats["delivery_hours"] == pytest.approx(960.0)
    assert stats["avg_price_reason"] is None
    # The row mean the page used to show is a different, denser-half-biased number.
    assert round(float(frame["price_eur_mwh"].mean()), 2) == pytest.approx(61.43)


@pytest.mark.parametrize("freq", ["h", "30min", "15min"])
def test_uniform_grids_keep_the_row_mean(freq: str) -> None:
    from src.analytics import calculate_average_price

    frame = _uniform_sample(freq)
    stats = calculate_average_price(frame)
    assert stats["avg_price_eur_mwh"] == pytest.approx(float(frame["price_eur_mwh"].mean()))
    assert stats["avg_price_eur_mwh"] == pytest.approx(32.50)
    assert stats["delivery_hours"] == pytest.approx(960.0)


@pytest.mark.parametrize("freq", ["2h", "24h"])
def test_regular_sparse_timestamps_do_not_invent_long_delivery_products(freq: str) -> None:
    from src.analytics import average_price_basis, calculate_average_price

    frame = _frame(
        pd.date_range("2026-01-01", periods=3, freq=freq, tz="UTC"),
        [10.0, 100.0, 10.0],
    )
    stats = calculate_average_price(frame)
    assert np.isnan(stats["avg_price_eur_mwh"])
    assert np.isnan(stats["covered_hours"])
    assert np.isnan(stats["delivery_hours"])
    assert "not a supported native" in stats["avg_price_reason"]
    assert average_price_basis(stats) is None


def test_numeric_row_index_is_not_interpreted_as_delivery_nanoseconds() -> None:
    from src.analytics import average_price_basis, calculate_average_price

    frame = pd.DataFrame({"price_eur_mwh": [10.0, 100.0, 10.0]})
    stats = calculate_average_price(frame)
    assert np.isnan(stats["avg_price_eur_mwh"])
    assert np.isnan(stats["covered_hours"])
    assert np.isnan(stats["delivery_hours"])
    assert stats["avg_price_reason"] == "the price index is not a timestamp index"
    assert average_price_basis(stats) is None


@pytest.mark.parametrize(
    "tz,hourly,quarter_hours,expected",
    [
        # FI civil midnight is 21:00Z: 49 hourly products, then 23 quarter-hour hours.
        ("Europe/Helsinki", 49, 23, (49 * 10 + 23 * 100) / 72),
        # PT civil midnight is 23:00Z: 47 hourly products, then 25 quarter-hour hours.
        ("Europe/Lisbon", 47, 25, (47 * 10 + 25 * 100) / 72),
    ],
)
def test_cutover_is_the_shared_market_instant_not_the_civil_midnight(
    tz: str, hourly: int, quarter_hours: int, expected: float,
) -> None:
    from src.analytics import calculate_average_price

    frame = _civil_cutover_sample(tz)
    assert int((frame["price_eur_mwh"] == 10.0).sum()) == hourly
    assert int((frame["price_eur_mwh"] == 100.0).sum()) == quarter_hours * 4
    stats = calculate_average_price(frame)
    assert stats["avg_price_eur_mwh"] == pytest.approx(expected)
    assert stats["delivery_hours"] == pytest.approx(72.0)


@pytest.mark.parametrize("day,hours", [("2026-03-29", 23.0), ("2026-10-25", 25.0)])
def test_dst_days_count_their_physical_hours(day: str, hours: float) -> None:
    from src.analytics import calculate_average_price

    start = pd.Timestamp(day, tz="Europe/Berlin")
    end = pd.Timestamp(pd.Timestamp(day) + pd.Timedelta(days=1), tz="Europe/Berlin")
    index = pd.date_range(
        start.tz_convert("UTC"), end.tz_convert("UTC"), inclusive="left", freq="15min",
    )
    # First local hour at 100, the remaining physical hours at 10.
    values = np.where(np.arange(len(index)) < 4, 100.0, 10.0)
    stats = calculate_average_price(_frame(index, values))
    assert len(index) == hours * 4
    assert stats["delivery_hours"] == pytest.approx(hours)
    assert stats["avg_price_eur_mwh"] == pytest.approx((100.0 + 10.0 * (hours - 1)) / hours)


def test_non_finite_prices_are_excluded_and_disclosed_not_counted_as_zero() -> None:
    from src.analytics import average_price_basis, calculate_average_price

    index = pd.date_range("2026-01-01", periods=6, freq="h", tz="UTC")
    stats = calculate_average_price(
        _frame(index, [10.0, np.nan, np.inf, -np.inf, 20.0, 30.0]),
    )
    assert stats["avg_price_eur_mwh"] == pytest.approx(20.0)
    assert stats["covered_hours"] == pytest.approx(3.0)
    assert stats["delivery_hours"] == pytest.approx(6.0)
    basis = average_price_basis(stats)
    assert "3.00 of 6.00 delivery hours (50.0%)" in basis
    assert "excluded, not counted as zero" in basis


def test_full_coverage_basis_does_not_claim_an_exclusion() -> None:
    from src.analytics import average_price_basis, calculate_average_price

    basis = average_price_basis(calculate_average_price(_cutover_sample()))
    assert basis == (
        "Duration-weighted: finite prices cover 960.00 of 960.00 delivery hours (100.0%)."
    )


def test_no_finite_price_is_unavailable_with_known_delivery_hours() -> None:
    from src.analytics import average_price_basis, calculate_average_price

    index = pd.date_range("2026-01-01", periods=4, freq="15min", tz="UTC")
    stats = calculate_average_price(_frame(index, [np.nan, np.inf, np.nan, -np.inf]))
    assert np.isnan(stats["avg_price_eur_mwh"])
    assert stats["covered_hours"] == 0.0
    assert stats["delivery_hours"] == pytest.approx(1.0)
    assert stats["avg_price_reason"] == "no delivery interval carries a finite price"
    assert average_price_basis(stats) is None


@pytest.mark.parametrize(
    "sample,needle",
    [
        ("gap", _GAP_REASON),
        ("unregistered_cadence_change", _GAP_REASON),
        ("duplicate", "repeats the same delivery timestamp"),
        ("unsorted", "not sorted into delivery order"),
        ("nat", "unset (NaT) timestamps"),
        ("single", "single delivery timestamp"),
        ("empty", "no price observations"),
    ],
)
def test_unverifiable_duration_is_unavailable_not_guessed(sample: str, needle: str) -> None:
    from src.analytics import average_price_basis, calculate_average_price

    if sample == "gap":
        frame = _gap_sample()
    elif sample == "unregistered_cadence_change":
        # Hour-to-quarter change on a day that is not the registered cutover.
        before = pd.date_range("2026-01-01", periods=24, freq="h", tz="UTC")
        after = pd.date_range("2026-01-02", periods=96, freq="15min", tz="UTC")
        frame = _frame(before.append(after), np.r_[np.full(24, 10.0), np.full(96, 100.0)])
    elif sample == "single":
        # One price alone cannot show whether it covered an hour or a quarter.
        frame = _frame(pd.DatetimeIndex(["2026-01-01T00:00Z"]), [42.0])
    elif sample == "empty":
        frame = pd.DataFrame(
            {"price_eur_mwh": pd.Series(dtype=float)},
            index=pd.DatetimeIndex([], tz="UTC", name="timestamp"),
        )
    else:
        frame = _defective_sample(sample)
    stats = calculate_average_price(frame)
    assert np.isnan(stats["avg_price_eur_mwh"])
    assert np.isnan(stats["covered_hours"])
    assert needle in stats["avg_price_reason"]
    assert average_price_basis(stats) is None


# ── Market page ──────────────────────────────────────────────────────────────

def _market_app(sample: str) -> None:
    from src.analytics import (
        calculate_daily_spreads,
        calculate_negative_price_hours,
        calculate_spread_percentiles,
    )
    from src.pages.market_overview import render
    from tests.test_step3c_average_price import _cutover_sample, _gap_sample, _uniform_sample

    frame = {
        "cutover": _cutover_sample,
        "gap": _gap_sample,
        "hourly": lambda: _uniform_sample("h"),
    }[sample]()
    daily = calculate_daily_spreads(frame)
    render(
        "DE_LU", frame, daily, calculate_spread_percentiles(daily),
        calculate_negative_price_hours(frame), 1, "UTC", "plotly_dark", {},
    )


def _avg_price_metric(app):
    return next(m for m in app.metric if m.label.startswith("Avg Price"))


def _run_market(sample: str):
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_function(_market_app, args=(sample,)).run(timeout=60)
    assert not app.exception
    return app


def test_page_shows_the_duration_weighted_average_and_its_coverage() -> None:
    app = _run_market("cutover")
    assert "32.50" in _avg_price_metric(app).value
    assert "61.43" not in _avg_price_metric(app).value
    assert any(
        "960.00 of 960.00 delivery hours (100.0%)" in c.value for c in app.caption
    )


def test_page_shows_na_and_the_reason_when_duration_is_unverifiable() -> None:
    app = _run_market("gap")
    assert _avg_price_metric(app).value == "n/a"
    assert any(
        "Avg Price is unavailable because" in c.value and _GAP_REASON in c.value
        for c in app.caption
    )


def test_page_keeps_a_uniform_hourly_average() -> None:
    """Compatibility control: a uniform grid shows the same value as before."""
    app = _run_market("hourly")
    assert "32.50" in _avg_price_metric(app).value


# ── Excel / PDF summaries ────────────────────────────────────────────────────

def _summary_pairs(data: bytes) -> dict[str, object]:
    ws = load_workbook(BytesIO(data))["Summary"]
    return {
        row[0].value: row[1].value
        for row in ws.iter_rows(min_col=1, max_col=2)
        if row[0].value
    }


def test_excel_summary_uses_the_duration_weighted_average() -> None:
    data = export_to_bytes(**_export_args(_cutover_sample()))
    pairs = _summary_pairs(data)
    assert pairs["Avg Price (EUR/MWh)"] == pytest.approx(32.50)
    assert "960.00 of 960.00 delivery hours" in pairs["Avg Price Basis"]
    assert "Avg Price Unavailable Because" not in pairs
    # The companion median remains an equal-row statistic, with a visible basis.
    assert pairs["Median Price (row-based, EUR/MWh)"] == 100.0
    assert isinstance(pairs["Median Price (row-based, EUR/MWh)"], (int, float))
    ws = load_workbook(BytesIO(data))["Summary"]
    median_label = next(
        row[0] for row in ws.iter_rows(min_col=1, max_col=2)
        if row[0].value == "Median Price (row-based, EUR/MWh)"
    )
    median_cell = ws.cell(row=median_label.row, column=2)
    assert median_cell.data_type == "n"
    assert median_cell.number_format == "#,##0.00"


def test_excel_summary_shows_na_and_reason_when_duration_is_unverifiable() -> None:
    pairs = _summary_pairs(export_to_bytes(**_export_args(_gap_sample())))
    assert pairs["Avg Price (EUR/MWh)"] == "n/a"
    assert _GAP_REASON in pairs["Avg Price Unavailable Because"]
    assert "Avg Price Basis" not in pairs
    for value in pairs.values():
        assert not (isinstance(value, float) and not np.isfinite(value))
        assert str(value).strip().lower() != "nan"


@pytest.mark.parametrize("freq", ["h", "30min", "15min"])
def test_excel_summary_keeps_uniform_averages(freq: str) -> None:
    """Compatibility control: the numeric cell and its value are unchanged."""
    pairs = _summary_pairs(export_to_bytes(**_export_args(_uniform_sample(freq))))
    assert pairs["Avg Price (EUR/MWh)"] == pytest.approx(32.50)
    assert isinstance(pairs["Avg Price (EUR/MWh)"], float)


def test_pdf_summary_uses_the_duration_weighted_average() -> None:
    text = _pdf_text(export_to_pdf_bytes(**_export_args(_cutover_sample())))
    assert "Avg Price (EUR/MWh) 32.50" in text
    assert "61.43" not in text
    assert "960.00 of 960.00 delivery hours (100.0%)" in text
    assert "Median Price (row-based, EUR/MWh) 100.00" in text


def test_pdf_summary_shows_na_and_reason_when_duration_is_unverifiable() -> None:
    text = _pdf_text(export_to_pdf_bytes(**_export_args(_gap_sample())))
    assert "Avg Price (EUR/MWh) n/a" in text
    assert "Avg Price Unavailable Because" in text
    assert "nan" not in text.lower()


def test_pdf_summary_keeps_a_uniform_hourly_average() -> None:
    """Compatibility control."""
    text = _pdf_text(export_to_pdf_bytes(**_export_args(_uniform_sample("h"))))
    assert "Avg Price (EUR/MWh) 32.50" in text


def test_pdf_disclosure_rows_stay_inside_the_page() -> None:
    """A long basis sentence wraps in its column instead of leaving the page."""
    from fpdf import FPDF

    from src.analytics import average_price_basis, calculate_average_price

    frame = _cutover_sample()
    frame.iloc[5, 0] = np.nan  # adds the longer "excluded, not counted as zero" clause
    stats = calculate_average_price(frame)
    sentence = stats["avg_price_reason"] or average_price_basis(stats)
    chunks = _pdf_chunks(export_to_pdf_bytes(**_export_args(frame)))
    ruler = FPDF(orientation="L", unit="mm", format="A4")
    ruler.add_page()
    available = ruler.w - ruler.r_margin - (ruler.l_margin + 100)
    ruler.set_font("Helvetica", "", 10)
    # The premise: the sentence is wider than its value column.
    assert ruler.get_string_width(sentence) > available
    start = next(i for i, chunk in enumerate(chunks) if chunk and sentence.startswith(chunk))
    lines: list[str] = []
    while len(" ".join(lines)) < len(sentence):
        lines.append(chunks[start + len(lines)])
    assert " ".join(lines) == sentence
    assert len(lines) >= 2
    for line in lines:
        assert ruler.get_string_width(line) <= available, line


# ── Zone comparison ──────────────────────────────────────────────────────────

def test_zone_comparison_uses_the_shared_average_and_discloses_coverage() -> None:
    comp = compare_zones({"DE_LU": _cutover_sample(), "FR": _gap_sample()})
    rows = comp.set_index("zone")
    assert rows.loc["DE_LU", "avg_price"] == pytest.approx(32.50)
    assert rows.loc["DE_LU", "avg_price_coverage_pct"] == pytest.approx(100.0)
    assert rows.loc["DE_LU", "avg_price_unavailable_reason"] is None
    assert np.isnan(rows.loc["FR", "avg_price"])
    assert np.isnan(rows.loc["FR", "avg_price_coverage_pct"])
    assert _GAP_REASON in rows.loc["FR", "avg_price_unavailable_reason"]


def test_zone_comparison_keeps_a_uniform_hourly_average() -> None:
    """Compatibility control."""
    comp = compare_zones({"DE_LU": _uniform_sample("h")})
    assert comp["avg_price"].iloc[0] == pytest.approx(32.50)


def test_comparison_export_writes_na_beside_the_reason() -> None:
    comp = compare_zones({"DE_LU": _cutover_sample(), "FR": _gap_sample()})
    ws = load_workbook(BytesIO(export_comparison_to_bytes(comp)))["Zone Comparison"]
    header = [cell.value for cell in ws[1]]
    rows = {
        row[header.index("Zone")]: row
        for row in ws.iter_rows(min_row=2, values_only=True)
    }
    avg = header.index("Avg Price (EUR/MWh)")
    coverage = header.index("Avg Price Coverage %")
    reason = header.index("Avg Price Unavailable Because")
    assert rows["DE_LU"][avg] == pytest.approx(32.50)
    assert rows["DE_LU"][coverage] == pytest.approx(1.0)
    assert rows["DE_LU"][reason] is None
    assert rows["FR"][avg] == "n/a"
    assert rows["FR"][coverage] == "n/a"
    assert _GAP_REASON in rows["FR"][reason]
    cells = {
        row[header.index("Zone")].value: row
        for row in ws.iter_rows(min_row=2)
    }
    reason_cell = cells["FR"][reason]
    assert reason_cell.value == comp.set_index("zone").loc["FR", "avg_price_unavailable_reason"]
    assert reason_cell.alignment.wrap_text is True
    assert ws.row_dimensions[reason_cell.row].height > ws.sheet_format.defaultRowHeight
    assert cells["DE_LU"][avg].data_type == "n"
    assert cells["DE_LU"][coverage].data_type == "n"
    assert cells["DE_LU"][coverage].number_format == "0.0%"
    std = header.index("Std Dev (row-based)")
    assert cells["DE_LU"][std].data_type == "n"
    assert cells["DE_LU"][std].number_format == "#,##0.00"
    assert rows["DE_LU"][std] == pytest.approx(44.55)


def _zone_comparison_app() -> None:
    from src.pages.zone_comparison import render
    from tests.test_step3c_average_price import _cutover_sample, _gap_sample

    render(
        {"DE_LU": _cutover_sample(), "FR": _gap_sample()},
        duration_hours=1, capture_rate=0.7, efficiency=0.88, power_mw=1.0,
        use_lp_dispatch=False, capex_eur_kwh=0.0, chart_template="plotly_dark",
    )


def test_zone_comparison_page_names_the_unavailable_zone_and_reason() -> None:
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_function(_zone_comparison_app).run(timeout=120)
    assert not app.exception
    assert any(
        c.value.startswith("Avg Price for FR: n/a — ") and _GAP_REASON in c.value
        for c in app.caption
    )
    assert not any("Avg Price for DE_LU" in c.value for c in app.caption)
    columns = json.loads(app.dataframe[0].proto.columns)
    assert columns["std_price"]["label"] == "Std Dev (row-based)"
    assert "Each row has equal weight" in columns["std_price"]["help"]
