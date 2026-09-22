"""Content identity and real-panel regressions for cached frontier results.

The signature adapter intentionally also runs on the pre-fix baseline: old
code collects normally, compatibility controls pass, and missing content
identity / export snapshots fail assertions rather than imports.
"""

from __future__ import annotations

import inspect
import json
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import src.cycle_frontier as cycle_frontier
import src.degradation as degradation
import src.dispatch as dispatch
import src.pages.simulation_cockpit as cockpit


def _prices() -> pd.DataFrame:
    day = [10.0] * 6 + [100.0] * 3 + [10.0] * 6 + [100.0] * 3 + [50.0] * 6
    index = pd.date_range("2026-03-02", periods=96, freq="h", tz="UTC", name="timestamp")
    return pd.DataFrame({"price_eur_mwh": day * 4}, index=index)


def _selected_dates() -> list:
    # Three selected days; March 4 is loaded but unselected. Replacing the
    # interior day therefore preserves the old (first, last, count) key.
    return [pd.Timestamp(day).date() for day in ("2026-03-02", "2026-03-03", "2026-03-05")]


def _identity_kwargs() -> dict:
    return {
        "primary_zone": "DE_LU",
        "primary_df": _prices(),
        "caps": [1.0, None],
        "cycle_life": 6000.0,
        "sweep_dates": _selected_dates(),
        "zone_tz": "UTC",
        "power_mw": 1.0,
        "duration_hours": 1.0,
        "efficiency": 0.9,
        "capex_eur_kwh": 150.0,
    }


def _identity(**overrides) -> tuple:
    kwargs = {**_identity_kwargs(), **overrides}
    if "primary_df" not in inspect.signature(cockpit._frontier_fingerprint).parameters:
        kwargs.pop("primary_df")
    return cockpit._frontier_fingerprint(**kwargs)


def _legacy_identity() -> tuple:
    kwargs = _identity_kwargs()
    days = kwargs["sweep_dates"]
    return (
        "DE_LU", (-1.0, 1.0), 6000.0,
        (str(days[0]), str(days[-1]), len(days)), "UTC", 1.0, 1.0, 0.9, 150.0,
    )


def _corrected_prices(kind: str) -> pd.DataFrame:
    frame = _prices()
    if kind == "selected_price":
        frame.iloc[7, 0] = 1200.0
    elif kind == "unselected_price":
        frame.iloc[55, 0] = 1200.0
    elif kind == "interior_index":
        index = frame.index.to_list()
        index[7] += pd.Timedelta(minutes=15)
        frame.index = pd.DatetimeIndex(index, name=frame.index.name)
    elif kind == "dtype":
        frame = frame.astype({"price_eur_mwh": "int64"})
    elif kind == "timezone":
        frame.index = frame.index.tz_convert("Europe/Berlin")
    elif kind == "index_name":
        frame.index.name = "corrected_timestamp"
    elif kind == "column_name":
        frame = frame.rename(columns={"price_eur_mwh": "corrected_price"})
    elif kind == "row_order":
        frame = frame.iloc[[1, 0, *range(2, len(frame))]]
    elif kind == "extra_column":
        frame["source_revision"] = "revised"
    else:
        raise AssertionError(f"Unknown correction: {kind}")
    return frame


def test_primary_frame_is_a_required_keyword_input() -> None:
    parameter = inspect.signature(cockpit._frontier_fingerprint).parameters.get("primary_df")
    assert parameter is not None, "Frontier identity must require the actual price frame"
    assert parameter.kind is inspect.Parameter.KEYWORD_ONLY
    assert parameter.default is inspect.Parameter.empty


def test_equal_frame_copies_and_reordered_caps_match_compat() -> None:
    identity = _identity()
    assert isinstance(identity, tuple)  # The contracted-floor fingerprint consumes a tuple.
    assert identity == _identity(primary_df=_prices().copy(deep=True), caps=[None, 1.0])


def test_disabled_liquidity_ignores_its_inactive_parameters_compat() -> None:
    assert _identity() == _identity(
        liquidity_enabled=False, zone_da_volume_mw=123.0, max_participation_share=0.75,
    )


def test_uncapped_sentinel_cannot_collide_with_zero_cap_compat() -> None:
    assert _identity(caps=[None]) != _identity(caps=[0.0])


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("primary_zone", "AT"),
        ("caps", [0.5, None]),
        ("cycle_life", 3000.0),
        ("sweep_dates", _selected_dates()[:2]),
        ("zone_tz", "Europe/Berlin"),
        ("power_mw", 2.0),
        ("duration_hours", 2.0),
        ("efficiency", 0.8),
        ("capex_eur_kwh", 100.0),
        ("liquidity_enabled", True),
    ],
)
def test_existing_active_inputs_still_invalidate_compat(field: str, value) -> None:
    assert _identity() != _identity(**{field: value})


@pytest.mark.parametrize(
    ("field", "value"),
    [("zone_da_volume_mw", 20.0), ("max_participation_share", 0.2)],
)
def test_enabled_liquidity_values_still_invalidate_compat(field: str, value: float) -> None:
    kwargs = {"liquidity_enabled": True, "zone_da_volume_mw": 10.0, "max_participation_share": 0.1}
    assert _identity(**kwargs) != _identity(**{**kwargs, field: value})


@pytest.mark.parametrize(
    "kind",
    [
        "selected_price", "unselected_price", "interior_index", "dtype", "timezone",
        "index_name", "column_name", "row_order", "extra_column",
    ],
)
def test_full_frame_content_and_layout_are_part_of_identity(kind: str) -> None:
    corrected = _corrected_prices(kind)
    assert len(corrected) == len(_prices())
    assert _identity() != _identity(primary_df=corrected)


@pytest.mark.parametrize("change", ["interior_replacement", "order"])
def test_exact_ordered_selected_dates_are_part_of_identity(change: str) -> None:
    original = _selected_dates()
    days = original.copy()
    if change == "interior_replacement":
        days[1] = pd.Timestamp("2026-03-04").date()
    else:
        original.insert(2, pd.Timestamp("2026-03-04").date())
        days = [original[0], original[2], original[1], original[3]]
    assert (days[0], days[-1], len(days)) == (original[0], original[-1], len(original))
    assert _identity(sweep_dates=original) != _identity(sweep_dates=days)


def test_panel_version_participates_in_identity(monkeypatch: pytest.MonkeyPatch) -> None:
    version = getattr(cockpit, "_FRONTIER_PANEL_ID", None)
    assert version, "A panel version must invalidate results after semantic code changes"
    original = _identity()
    monkeypatch.setattr(cockpit, "_FRONTIER_PANEL_ID", version + "/test-next")
    assert _identity() != original


@pytest.mark.parametrize(
    ("module", "name"),
    [
        (dispatch, "DISPATCH_VOM_COST_EUR_MWH"),
        (cycle_frontier, "DAYS_PER_YEAR"),
        (degradation, "DAYS_PER_YEAR"),
        (cycle_frontier, "NET_TOL_EUR_PER_MW_YR"),
    ],
)
def test_consumed_solver_constants_participate_in_identity(
    monkeypatch: pytest.MonkeyPatch, module, name: str,
) -> None:
    original = _identity()
    monkeypatch.setattr(module, name, getattr(module, name) + 1.0)
    assert _identity() != original


def test_pre_content_identity_cannot_match_a_current_result() -> None:
    assert _identity() != _legacy_identity()


def _frontier_identity_app() -> None:
    """Real frontier and floor panels, supplied only small deterministic inputs."""
    import streamlit as st

    from src.pages.simulation_cockpit import (
        _render_contracted_floor_section,
        _render_cycle_frontier_section,
    )

    light = st.checkbox("Harness light theme", key="identity_light_theme")
    template = "plotly_white" if light else "plotly_dark"
    context = _render_cycle_frontier_section(
        primary_zone="DE_LU",
        primary_df=st.session_state["identity_prices"],
        dates=st.session_state["identity_dates"],
        zone_tz="UTC",
        power_mw=1.0,
        duration_hours=1,
        efficiency=0.9,
        capex_eur_kwh=150.0,
        chart_template=template,
        assumptions=st.session_state["identity_assumptions"],
    )
    st.session_state["identity_returned_context"] = context
    _render_contracted_floor_section(
        frontier_context=context,
        chart_template=template,
        assumptions=st.session_state["identity_assumptions"],
    )


def _elements(node, kind: str) -> list:
    found = []
    children = getattr(node, "children", None)
    if isinstance(children, dict):
        for child in children.values():
            if getattr(child, "type", "") == kind:
                found.append(child)
            found.extend(_elements(child, kind))
    return found


def _assert_current(app: AppTest, *, floor: bool = False) -> None:
    assert not app.exception
    assert app.session_state["identity_returned_context"] is not None
    labels = {metric.label for metric in app.metric}
    assert "Best cap (net of wear)" in labels
    assert ("Annual merchant net" in labels) is floor
    assert len(app.dataframe) >= (2 if floor else 1)
    downloads = _elements(app.main, "download_button")
    assert len(downloads) == (2 if floor else 1)
    assert any("frontier" in item.proto.label for item in downloads)


def _assert_stale(app: AppTest) -> None:
    assert not app.exception
    assert app.session_state["identity_returned_context"] is None
    assert any("Inputs changed" in item.value for item in app.info)
    assert any("Run a valid cycle-cap frontier" in item.value for item in app.info)
    assert not app.dataframe
    assert not _elements(app.main, "plotly_chart")
    assert not _elements(app.main, "download_button")
    labels = {metric.label for metric in app.metric}
    assert "Best cap (net of wear)" not in labels
    assert "Annual merchant net" not in labels


@pytest.fixture()
def instrumented_app(monkeypatch: pytest.MonkeyPatch):
    observed = {"compute": [], "solve": [], "exports": []}
    real_compute = cockpit.compute_cycle_cap_frontier
    real_solve = cycle_frontier.solve_daily_lp
    real_export = cockpit.cockpit_tables_to_excel

    def recording_compute(frame, **kwargs):
        # Instrument the real entry point; no canned frame or summary can
        # accidentally make a broken compute/cache/export path look correct.
        observed["compute"].append((frame.copy(deep=True), dict(kwargs)))
        return real_compute(frame, **kwargs)

    def recording_solve(*args, **kwargs):
        observed["solve"].append(dict(kwargs))
        return real_solve(*args, **kwargs)

    def recording_export(tables, *, assumptions=None):
        data = real_export(tables, assumptions=assumptions)
        observed["exports"].append((tuple(tables), data))
        return data

    monkeypatch.setattr(cockpit, "compute_cycle_cap_frontier", recording_compute)
    monkeypatch.setattr(cycle_frontier, "solve_daily_lp", recording_solve)
    monkeypatch.setattr(cockpit, "cockpit_tables_to_excel", recording_export)
    app = AppTest.from_function(_frontier_identity_app)
    app.session_state["identity_prices"] = _prices()
    app.session_state["identity_dates"] = _selected_dates()
    app.session_state["identity_assumptions"] = pd.DataFrame([
        {
            "parameter": "Harness provenance", "value": "original-run", "unit": "",
            "source": "Harness sidebar", "affects": "Export audit",
        },
    ])
    app.session_state["cycle_frontier_caps"] = ["1", "uncapped"]
    app.run(timeout=30)
    assert not app.exception
    assert not observed["compute"]
    assert not observed["solve"]
    assert not observed["exports"]
    assert not _elements(app.main, "download_button")
    return app, observed


def _run_frontier(app: AppTest) -> None:
    app.button(key="cycle_frontier_run").click().run(timeout=60)
    assert not app.exception


def _frontier_workbooks(observed: dict) -> list[bytes]:
    return [data for sheets, data in observed["exports"] if "Cycle-cap frontier" in sheets]


def _assumptions_from_workbook(data: bytes) -> pd.DataFrame:
    workbook = load_workbook(BytesIO(data), data_only=True)
    assert "Cycle-cap frontier" in workbook.sheetnames
    rows = list(workbook["Assumptions"].values)
    return pd.DataFrame(rows[1:], columns=rows[0]).set_index("parameter")


def test_real_frontier_reuses_result_on_rerun_and_theme_change_compat(instrumented_app) -> None:
    app, observed = instrumented_app
    _run_frontier(app)
    _assert_current(app)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    assert len(app.dataframe[0].value) == 2
    assert app.session_state["identity_returned_context"]["summary"]["valid_days"] == 3
    chart_before = json.loads(_elements(app.main, "plotly_chart")[0].proto.spec)
    original_table = app.dataframe[0].value.copy(deep=True)

    app.run(timeout=30)
    _assert_current(app)
    app.checkbox(key="identity_light_theme").check().run(timeout=30)
    _assert_current(app)
    chart_after = json.loads(_elements(app.main, "plotly_chart")[0].proto.spec)
    assert app.checkbox(key="identity_light_theme").value is True
    # Cockpit styling deliberately keeps charts dark for either global theme.
    assert chart_before["data"] == chart_after["data"]
    pd.testing.assert_frame_equal(app.dataframe[0].value, original_table)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.

    app.button(key="contracted_floor_run").click().run(timeout=30)
    _assert_current(app, floor=True)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    _run_frontier(app)
    _assert_current(app, floor=True)
    assert len(observed["compute"]) == 2
    assert len(observed["solve"]) == 12


@pytest.mark.parametrize(
    "correction", ["selected_price", "interior_index", "unselected_price", "interior_date"],
)
def test_stale_source_hides_frontier_floor_and_downloads_without_solving(
    instrumented_app, correction: str,
) -> None:
    app, observed = instrumented_app
    _run_frontier(app)
    app.button(key="contracted_floor_run").click().run(timeout=30)
    _assert_current(app, floor=True)
    original_table = app.dataframe[0].value.copy(deep=True)
    exports_before = len(observed["exports"])

    if correction == "interior_date":
        days = _selected_dates()
        days[1] = pd.Timestamp("2026-03-04").date()
        assert (days[0], days[-1], len(days)) == (
            _selected_dates()[0], _selected_dates()[-1], len(_selected_dates()),
        )
        app.session_state["identity_dates"] = days
    else:
        corrected = _corrected_prices(correction)
        assert corrected.shape == _prices().shape
        assert corrected.index[[0, -1]].equals(_prices().index[[0, -1]])
        app.session_state["identity_prices"] = corrected
    app.run(timeout=30)
    _assert_stale(app)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    assert len(observed["exports"]) == exports_before
    app.run(timeout=30)
    _assert_stale(app)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    assert len(observed["exports"]) == exports_before

    app.session_state["identity_prices"] = _prices()
    app.session_state["identity_dates"] = _selected_dates()
    app.run(timeout=30)
    # The frontier cache survives staleness. Existing floor behavior clears
    # its cache when its source context is None, so it needs its own Run.
    _assert_current(app)
    pd.testing.assert_frame_equal(app.dataframe[0].value, original_table)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    app.button(key="contracted_floor_run").click().run(timeout=30)
    _assert_current(app, floor=True)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    _run_frontier(app)
    _assert_current(app, floor=True)
    assert len(observed["compute"]) == 2
    assert len(observed["solve"]) == 12


def test_run_after_correction_recomputes_current_data_once(instrumented_app) -> None:
    app, observed = instrumented_app
    _run_frontier(app)
    old_result = app.session_state["identity_returned_context"]["frontier"].copy(deep=True)
    corrected = _corrected_prices("selected_price")
    app.session_state["identity_prices"] = corrected
    app.run(timeout=30)
    _assert_stale(app)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    _run_frontier(app)
    _assert_current(app)
    assert len(observed["compute"]) == 2
    assert len(observed["solve"]) == 12
    pd.testing.assert_frame_equal(observed["compute"][-1][0], corrected)
    assert observed["compute"][-1][1]["dates"] == _selected_dates()
    new_result = app.session_state["identity_returned_context"]["frontier"]
    assert new_result["gross_eur"].max() > old_result["gross_eur"].max()
    app.run(timeout=30)
    _assert_current(app)
    assert len(observed["compute"]) == 2
    assert len(observed["solve"]) == 12


def test_real_excel_download_keeps_run_assumptions_until_next_run(instrumented_app) -> None:
    app, observed = instrumented_app
    _run_frontier(app)
    initial = _assumptions_from_workbook(_frontier_workbooks(observed)[-1])
    assert initial.loc["Harness provenance", "value"] == "original-run"
    assert initial.loc["Cockpit capture haircut", "value"] == "Not applied (raw solver values)"
    assert initial.loc["Frontier capex basis", "value"] == "150"
    assert initial.loc["Frontier cycle life", "value"] == "6000"

    # Mutate the same caller-owned frame to catch accidental reference reuse.
    assumptions = app.session_state["identity_assumptions"]
    assumptions.loc[0, "value"] = "edited-after-run"
    app.checkbox(key="identity_light_theme").check().run(timeout=30)
    _assert_current(app)
    after_edit = _assumptions_from_workbook(_frontier_workbooks(observed)[-1])
    pd.testing.assert_frame_equal(after_edit, initial)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.

    _run_frontier(app)
    latest = _assumptions_from_workbook(_frontier_workbooks(observed)[-1])
    assert latest.loc["Harness provenance", "value"] == "edited-after-run"
    assert len(observed["compute"]) == 2
    assert len(observed["solve"]) == 12


def test_legacy_session_cache_is_hidden_until_a_new_run(instrumented_app) -> None:
    app, observed = instrumented_app
    _run_frontier(app)
    current = app.session_state["cycle_frontier_result"]
    app.session_state["cycle_frontier_result"] = {
        "fingerprint": _legacy_identity(),
        "frontier": current["frontier"],
        "summary": current["summary"],
        "liquidity": current["liquidity"],
    }
    app.run(timeout=30)
    _assert_stale(app)
    assert len(observed["compute"]) == 1
    assert len(observed["solve"]) == 6  # Three days times two caps.
    _run_frontier(app)
    _assert_current(app)
    assert len(observed["compute"]) == 2
    assert len(observed["solve"]) == 12


@pytest.mark.parametrize("window", ["empty_selection", "no_valid_days"])
def test_empty_frontier_keeps_download_and_floor_gated_compat(
    instrumented_app, window: str,
) -> None:
    app, observed = instrumented_app
    if window == "empty_selection":
        app.session_state["identity_dates"] = []
    else:
        frame = _prices()
        frame["price_eur_mwh"] = float("nan")
        app.session_state["identity_prices"] = frame
    _run_frontier(app)
    assert app.session_state["identity_returned_context"] is None
    assert any("No valid days" in item.value for item in app.warning)
    assert not app.dataframe
    assert not app.metric
    assert not _elements(app.main, "download_button")
    assert not observed["exports"]
    assert len(observed["compute"]) == 1
    assert not observed["solve"]
    app.run(timeout=30)
    assert not app.exception
    assert app.session_state["identity_returned_context"] is None
    assert not _elements(app.main, "download_button")
    assert not observed["exports"]
    assert len(observed["compute"]) == 1
    assert not observed["solve"]
