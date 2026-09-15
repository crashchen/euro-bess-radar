"""Step 3A: unavailable physical durations and defective price indices are disclosed.

These cases pin the user-visible contract: a duration that cannot be verified
renders as ``n/a`` plus a reason on every surface, the observed interval counts
survive, and a defective price index produces a visible diagnostic instead of a
bare page exception or a silently repaired chart.
"""

from __future__ import annotations

import contextlib
import re
import zlib
from io import BytesIO

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

# Symbols introduced by this step are imported inside the cases that use
# them, following tests/test_step2_views.py, so the module still collects and
# fails on behaviour against the pre-change baseline.
from src.analytics import (
    calculate_daily_spreads,
    calculate_negative_price_hours,
    calculate_spread_percentiles,
    estimate_annual_arbitrage_revenue,
)
from src.export import export_to_bytes, export_to_pdf_bytes

_UNVERIFIED = "could not be verified"


def _unverifiable_frame() -> pd.DataFrame:
    """A complete negative hourly day followed by a day with an unknown grid."""
    day_one = pd.date_range("2026-01-01", periods=24, freq="h", tz="UTC")
    day_two = pd.to_datetime(
        ["2026-01-02T00:00Z", "2026-01-02T01:00Z", "2026-01-02T03:00Z"],
    )
    index = day_one.append(pd.DatetimeIndex(day_two)).rename("timestamp")
    values = np.r_[np.full(24, -50.0), np.full(3, 60.0)]
    return pd.DataFrame({"price_eur_mwh": values}, index=index)


def _verifiable_frame(price: float = -50.0) -> pd.DataFrame:
    index = pd.date_range("2026-01-01", periods=48, freq="h", tz="UTC").rename("timestamp")
    return pd.DataFrame({"price_eur_mwh": np.full(48, price)}, index=index)


def _export_args(frame: pd.DataFrame) -> dict:
    daily = calculate_daily_spreads(frame)
    return {
        "zone": "DE_LU",
        "price_df": frame,
        "daily_spreads": daily,
        "monthly_spreads": pd.DataFrame(),
        "percentiles": calculate_spread_percentiles(daily),
        "revenue_estimate": estimate_annual_arbitrage_revenue(daily),
        "negative_stats": calculate_negative_price_hours(frame),
        "tz": "UTC",
    }


def _pdf_chunks(data: bytes) -> list[str]:
    """Return each literal show-text run of an fpdf2-generated PDF, in order."""
    chunks: list[str] = []
    for stream in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", data, re.S):
        raw = stream.group(1)
        with contextlib.suppress(zlib.error):
            raw = zlib.decompress(raw)
        for literal in re.finditer(rb"\((?:\\.|[^\\()])*\)", raw):
            chunks.append(
                re.sub(r"\\([()\\])", r"\1", literal.group(0)[1:-1].decode("latin-1")),
            )
    return chunks


def _pdf_text(data: bytes) -> str:
    """Extract the literal show-text strings from an fpdf2-generated PDF."""
    return " ".join(_pdf_chunks(data))


# ── Calculation contract ─────────────────────────────────────────────────────

def test_unverifiable_grid_reports_a_reason_and_keeps_interval_counts() -> None:
    from src.analytics import negative_price_hours_reason

    stats = calculate_negative_price_hours(_unverifiable_frame())
    assert np.isnan(stats["negative_hours"])
    assert np.isnan(stats["total_negative_hours"])
    assert "2026-01-02" in stats["negative_hours_reason"]
    assert _UNVERIFIED in stats["negative_hours_reason"]
    # Counts are observations, not durations: they must survive.
    assert stats["negative_intervals"] == 24
    assert stats["pct_negative"] == pytest.approx(round(100 * 24 / 27, 2))
    assert stats["avg_negative_price"] == -50.0
    assert stats["most_negative_price"] == -50.0
    assert negative_price_hours_reason(stats) == stats["negative_hours_reason"]


@pytest.mark.parametrize("price,hours", [(-50.0, 48.0), (50.0, 0.0)])
def test_verifiable_grid_has_no_reason(price: float, hours: float) -> None:
    from src.analytics import negative_price_hours_reason

    stats = calculate_negative_price_hours(_verifiable_frame(price))
    assert stats["negative_hours"] == hours
    assert stats["negative_hours_reason"] is None
    assert negative_price_hours_reason(stats) is None


# ── Excel surface ────────────────────────────────────────────────────────────

def _summary_pairs(data: bytes) -> dict[str, object]:
    ws = load_workbook(BytesIO(data))["Summary"]
    return {
        row[0].value: row[1].value
        for row in ws.iter_rows(min_col=1, max_col=2)
        if row[0].value
    }


def test_excel_summary_shows_na_and_reason_for_unverifiable_hours() -> None:
    pairs = _summary_pairs(export_to_bytes(**_export_args(_unverifiable_frame())))
    assert pairs["Negative Price Hours"] == "n/a"
    reason = pairs["Negative Price Hours Unavailable Because"]
    assert _UNVERIFIED in reason and "2026-01-02" in reason
    assert pairs["Negative Price Intervals"] == 24
    # No cell may carry a raw NaN or the literal spelling of one.
    for value in pairs.values():
        assert not (isinstance(value, float) and not np.isfinite(value))
        assert str(value).strip().lower() != "nan"


@pytest.mark.parametrize("price,shown", [(-50.0, 48.0), (50.0, 0.0)])
def test_excel_summary_keeps_verifiable_hours_and_omits_the_reason(
    price: float, shown: float,
) -> None:
    pairs = _summary_pairs(export_to_bytes(**_export_args(_verifiable_frame(price))))
    assert pairs["Negative Price Hours"] == shown
    assert "Negative Price Hours Unavailable Because" not in pairs


# ── PDF surface ──────────────────────────────────────────────────────────────

def test_pdf_summary_shows_na_and_reason_instead_of_a_literal_nan() -> None:
    text = _pdf_text(export_to_pdf_bytes(**_export_args(_unverifiable_frame())))
    assert "Negative Price Hours n/a" in text
    assert "Negative Price Hours Unavailable Because" in text
    assert _UNVERIFIED in text
    assert "2026-01-02" in text
    assert "nan" not in text.lower()
    assert "Negative Price Intervals 24" in text


def test_pdf_reason_row_fits_inside_the_page_margins() -> None:
    """The explanation is only useful if it is actually readable on the page."""
    from fpdf import FPDF

    data = export_to_pdf_bytes(**_export_args(_unverifiable_frame()))
    ruler = FPDF(orientation="L", unit="mm", format="A4")
    ruler.add_page()
    # The summary's value column starts after a 100 mm bold label column.
    available = ruler.w - ruler.r_margin - (ruler.l_margin + 100)
    ruler.set_font("Helvetica", "", 10)
    reason = calculate_negative_price_hours(_unverifiable_frame())["negative_hours_reason"]
    assert ruler.get_string_width(reason) <= available
    # Every rendered run must fit the widest column it can occupy.
    ruler.set_font("Helvetica", "B", 10)
    for chunk in _pdf_chunks(data):
        assert ruler.get_string_width(chunk) <= available, chunk


def test_pdf_summary_keeps_verifiable_hours() -> None:
    text = _pdf_text(export_to_pdf_bytes(**_export_args(_verifiable_frame())))
    assert "Negative Price Hours 48.0" in text
    assert "Unavailable Because" not in text


# ── Price index diagnosis ────────────────────────────────────────────────────

def _defective(defect: str) -> pd.DataFrame:
    index = pd.date_range("2026-01-01", periods=48, freq="h", tz="UTC")
    values = np.arange(48.0)
    if defect == "duplicate":
        order = [*range(48), 10]
        index, values = index[order], values[order]
        index = index.sort_values()
        values = np.sort(values)
    elif defect == "unsorted":
        order = [*range(24, 48), *range(24)]
        index, values = index[order], values[order]
    elif defect == "nat":
        index = pd.DatetimeIndex([*index[:47], pd.NaT])
    return pd.DataFrame({"price_eur_mwh": values}, index=index.rename("timestamp"))


@pytest.mark.parametrize(
    "defect,plottable,needle",
    [
        ("duplicate", True, "repeats the same delivery timestamp"),
        ("unsorted", False, "not sorted into delivery order"),
        ("nat", False, "unset (NaT) timestamps"),
    ],
)
def test_defective_index_is_classified(defect: str, plottable: bool, needle: str) -> None:
    from src.analytics import describe_price_index_issue, time_weighted_rolling_price_mean

    issue = describe_price_index_issue(_defective(defect).index)
    assert issue is not None
    assert issue.plottable is plottable
    assert needle in issue.reason
    # The calculation itself stays strict and still refuses the input.
    with pytest.raises(ValueError, match="unique increasing finite timestamps"):
        time_weighted_rolling_price_mean(_defective(defect)["price_eur_mwh"])


def test_clean_index_has_no_issue() -> None:
    from src.analytics import describe_price_index_issue

    assert describe_price_index_issue(_verifiable_frame().index) is None


# ── Real market page ─────────────────────────────────────────────────────────

def _market_app(defect: str) -> None:
    import streamlit as st

    from src.analytics import (
        calculate_daily_spreads,
        calculate_negative_price_hours,
        calculate_spread_percentiles,
    )
    from src.pages.market_overview import render
    from tests.test_step3a_display_contract import _defective, _verifiable_frame

    frame = _verifiable_frame(50.0) if defect == "clean" else _defective(defect)
    daily = calculate_daily_spreads(frame)
    figures: dict[str, object] = {"price_ts": "stale-figure-from-an-earlier-state"}
    render(
        "DE_LU", frame, daily, calculate_spread_percentiles(daily),
        calculate_negative_price_hours(frame), 1, "UTC", "plotly_dark", figures,
    )
    st.session_state["figure_keys"] = sorted(figures)
    price_ts = figures.get("price_ts")
    st.session_state["trace_names"] = (
        [trace.name for trace in price_ts.data] if hasattr(price_ts, "data") else None
    )


def _run_market_app(defect: str):
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_function(_market_app, args=(defect,)).run(timeout=60)
    assert not app.exception
    return app


def test_clean_index_still_renders_the_chart_and_moving_average() -> None:
    app = _run_market_app("clean")
    assert app.session_state["trace_names"] == ["Day-Ahead", "30-Day MA"]
    assert not app.error
    assert not app.warning
    assert any("30-Day MA uses the trailing 720 physical hours" in c.value for c in app.caption)


def test_repeated_timestamps_keep_the_raw_chart_but_drop_the_moving_average() -> None:
    app = _run_market_app("duplicate")
    assert app.session_state["trace_names"] == ["Day-Ahead"]
    assert "price_ts" in app.session_state["figure_keys"]
    assert any(
        "30-Day MA is unavailable" in w.value
        and "repeats the same delivery timestamp" in w.value
        for w in app.warning
    )


@pytest.mark.parametrize(
    "defect,needle",
    [("unsorted", "not sorted into delivery order"), ("nat", "unset (NaT) timestamps")],
)
def test_unplottable_index_reports_a_reason_and_clears_the_stale_export_figure(
    defect: str, needle: str,
) -> None:
    app = _run_market_app(defect)
    # The stale figure seeded above must not survive into the PDF export.
    assert "price_ts" not in app.session_state["figure_keys"]
    assert app.session_state["trace_names"] is None
    assert any(
        "Day-Ahead price chart and 30-Day MA are unavailable" in e.value and needle in e.value
        for e in app.error
    )


@pytest.mark.parametrize("defect", ["duplicate", "unsorted", "nat"])
def test_defective_rows_still_reach_the_excel_export_unchanged(defect: str) -> None:
    """Suppressing a chart must not suppress the underlying rows."""
    frame = _defective(defect)
    ws = load_workbook(BytesIO(export_to_bytes(**_export_args(frame))))["Hourly Prices"]
    exported = [row[1].value for row in ws.iter_rows(min_row=2, min_col=1, max_col=2)]
    assert exported == list(frame["price_eur_mwh"])


def _unverifiable_market_app() -> None:
    from src.analytics import (
        calculate_daily_spreads,
        calculate_negative_price_hours,
        calculate_spread_percentiles,
    )
    from src.pages.market_overview import render
    from tests.test_step3a_display_contract import _unverifiable_frame

    frame = _unverifiable_frame()
    daily = calculate_daily_spreads(frame)
    render(
        "DE_LU", frame, daily, calculate_spread_percentiles(daily),
        calculate_negative_price_hours(frame), 1, "UTC", "plotly_dark", {},
    )


def test_page_discloses_unverifiable_negative_hours_without_dropping_counts() -> None:
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_function(_unverifiable_market_app).run(timeout=60)
    assert not app.exception
    assert any(m.value == "n/a" for m in app.metric)
    assert any(
        "Negative-price hours are unavailable because" in c.value
        and _UNVERIFIED in c.value
        and "24 observed negative interval(s)" in c.value
        for c in app.caption
    )
