"""Production page/export disclosure; prior economic values remain controls."""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from src.export import export_to_bytes, export_to_pdf_bytes, project_case_to_excel
from tests.test_step3a_display_contract import _export_args, _pdf_text, _verifiable_frame
from tests.test_step3d_project_case_disclosure import synthetic_reserve_result


def _pairs(sheet):
    return {r[0].value: r[1] for r in sheet.iter_rows(min_col=1, max_col=2) if r[0].value}


def _joint_args():
    args = _export_args(_verifiable_frame(40.0))
    args['revenue_estimate'].update({
        'joint_cooptimized_total_eur': 159614.25,
        'joint_cooptimized_avg_reserve_fraction': .95,
    })
    return args


@pytest.mark.parametrize('zone,product', [('DE_LU', 'FCR'), ('FI', 'FCR-N')])
def test_joint_summary_and_pdf_share_the_bound_basis(zone, product):
    from src.settlement_disclosure import screening_capacity_settlement_basis

    args = _joint_args()
    args['zone'] = zone
    basis = screening_capacity_settlement_basis(zone, product)
    args['revenue_estimate']['joint_capacity_settlement_basis'] = basis
    summary = load_workbook(BytesIO(export_to_bytes(**args)))['Summary']
    pairs = _pairs(summary)
    cell = pairs['Joint MILP Capacity Settlement Basis']
    assert cell.value == basis
    assert cell.alignment.wrap_text
    text = _pdf_text(export_to_pdf_bytes(**args))
    assert ' '.join(basis.split()) in text
    assert f'{zone} / {product}' in text
    assert pairs['Joint MILP Co-optimized Total (EUR)'].value == 159614.25
    assert pairs['Joint MILP Co-optimized Total (EUR)'].data_type == 'n'
    assert '159,614' in text
    if zone == 'FI':
        assert 'nominal 4h capacity blocks' not in text


def test_legacy_joint_export_explicitly_discloses_missing_product_identity():
    args = _joint_args()
    pairs = _pairs(load_workbook(BytesIO(export_to_bytes(**args)))['Summary'])
    basis = pairs['Joint MILP Capacity Settlement Basis'].value
    assert 'product identity unavailable' in basis
    assert 'physical' in basis.lower()
    assert ' '.join(basis.split()) in _pdf_text(export_to_pdf_bytes(**args))


def test_da_only_exports_do_not_claim_capacity_settlement():
    args = _export_args(_verifiable_frame(40.0))
    pairs = _pairs(load_workbook(BytesIO(export_to_bytes(**args)))['Summary'])
    assert 'Joint MILP Capacity Settlement Basis' not in pairs
    assert 'capacity settlement' not in _pdf_text(export_to_pdf_bytes(**args)).lower()


def test_project_case_excel_names_basis_but_keeps_fingerprinted_payload_exact():
    from src.settlement_disclosure import project_case_capacity_settlement_basis

    result = synthetic_reserve_result()
    before = result.to_payload()
    book = load_workbook(BytesIO(project_case_to_excel(result)))
    pairs = _pairs(book['Project Case NPVs'])
    cell = pairs['Reserve Capacity Settlement Basis']
    assert cell.value == project_case_capacity_settlement_basis(result)
    assert 'DE_LU / FCR [symmetric]' in cell.value
    assert 'nominal 4h' in cell.value
    assert cell.alignment.wrap_text
    assert pairs['Project Case Input Fingerprint'].value == result.input_fingerprint
    assert result.to_payload() == before
    # The original assumptions/provenance wire tree is not extended by display text.
    paths = [r[0].value for r in book['Assumptions & Provenance'].iter_rows(min_row=2)]
    assert not any('disclosure' in str(p) or 'capacity_settlement_basis' in str(p) for p in paths)
    assert book['Project Case NPVs']['B6'].data_type == 'n'


def test_appended_project_case_workbook_gets_same_run_bound_disclosure():
    from src.settlement_disclosure import project_case_capacity_settlement_basis

    result = synthetic_reserve_result()
    args = _export_args(_verifiable_frame(40.0))
    args['zone'] = 'FI'  # Ambient Market Report cannot re-label the recorded DE case.
    book = load_workbook(BytesIO(export_to_bytes(**args, project_case_result=result)))
    pairs = _pairs(book['Project Case NPVs'])
    assert pairs['Reserve Capacity Settlement Basis'].value == project_case_capacity_settlement_basis(result)
    assert 'Joint MILP Capacity Settlement Basis' not in _pairs(book['Summary'])


def _revenue_app(with_capacity=True, unavailable=False):
    import datetime as dt
    from unittest.mock import patch

    import pandas as pd
    import streamlit as st

    from src.ancillary import _build_standard_frame
    from src.pages import revenue_estimation as page
    from tests.test_step3a_display_contract import _export_args

    day = dt.date(2026, 3, 29)
    lo = pd.Timestamp(day, tz='Europe/Berlin')
    hi = lo + pd.DateOffset(days=1)
    index = pd.date_range(lo, hi, freq='15min', inclusive='left').tz_convert('UTC')
    frame = pd.DataFrame({'price_eur_mwh': 40.0}, index=index)
    args = _export_args(frame)
    if with_capacity:
        # A real aggregate input: two capacity products and an energy-only product.
        st.session_state['ancillary_df'] = pd.concat([
            _build_standard_frame(index, 'FCR', 'DE_LU', capacity=20.),
            _build_standard_frame(index, 'aFRR Up', 'DE_LU', capacity=10.),
            _build_standard_frame(index, 'mFRR Up', 'DE_LU', energy=1.),
        ])
        st.session_state['ancillary_zone'] = 'DE_LU'
        st.session_state['ancillary_dates'] = (str(day), str(day))
    export = args['revenue_estimate'].copy()
    # Skip unrelated sub-panels, preserving the real ancillary aggregation,
    # joint solver and surrounding production page/export integration.
    with patch.object(page, '_render_intraday_uplift_section'), \
         patch.object(page, '_render_sensitivity_table'), \
         patch.object(page, '_render_revenue_risk_analysis'), \
         patch.object(page, 'render_project_case_panel'):
        if unavailable:
            failed = pd.DataFrame()
            failed.attrs.update(model_available=False, excluded_days_due_to_solver_failure=1)
            solver_context = patch.object(page, 'solve_joint_capacity_batch', return_value=failed)
        else:
            from contextlib import nullcontext
            solver_context = nullcontext()
        with solver_context:
            page.render(
                primary_zone='DE_LU', primary_df=frame,
                daily_spreads=args['daily_spreads'], monthly_spreads=args['monthly_spreads'],
                percentiles=args['percentiles'], revenue=args['revenue_estimate'],
                start_date=day, end_date=day, power_mw=1., duration_hours=1,
                efficiency=1., capture_rate=1., capex_eur_kwh=0., use_lp_dispatch=False,
                zone_tz='Europe/Berlin', chart_template='plotly_dark', report_figures={},
                export_revenue=export,
            )
    st.session_state['test_export'] = export


def test_revenue_page_binds_actual_capacity_products_to_ui_and_exports():
    app = AppTest.from_function(_revenue_app).run(timeout=30)
    assert not app.exception
    export = app.session_state['test_export']
    basis = export['joint_capacity_settlement_basis']
    assert 'DE_LU' in basis and 'aggregate capacity: FCR, aFRR Up' in basis
    assert 'mFRR Up' not in basis  # Energy-only contribution is not a capacity product.
    assert basis in [c.value for c in app.caption]


def test_revenue_joint_cash_is_unchanged_on_spring_dst():
    app = AppTest.from_function(_revenue_app).run(timeout=30)
    assert not app.exception
    export = app.session_state['test_export']
    assert export['joint_cooptimized_capacity_eur'] == pytest.approx(30 * .95 * 23 * 365.25)
    assert export['joint_cooptimized_da_eur'] == 0.0


@pytest.mark.parametrize('capacity,failed', [(False, False), (True, True)])
def test_no_joint_result_has_no_settlement_disclosure(capacity, failed):
    app = AppTest.from_function(_revenue_app, args=(capacity, failed)).run(timeout=30)
    assert not app.exception
    assert 'joint_capacity_settlement_basis' not in app.session_state['test_export']
    assert not any('physical-hour screening' in c.value for c in app.caption)
