"""Saved cockpit workbooks describe the solver and wear basis actually used."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.assumptions import build_assumptions_table
from src.export import cockpit_tables_to_excel
from src.pages import simulation_cockpit as cockpit


def _global_assumptions() -> pd.DataFrame:
    return build_assumptions_table(
        power_mw=1.0, duration_hours=2.0, efficiency=0.88,
        capture_rate=0.7, capex_eur_kwh=150.0, use_lp_dispatch=False,
    )


def _workbook_rows(data: bytes, sheet: str) -> tuple[object, dict[str, dict]]:
    workbook = load_workbook(BytesIO(data), data_only=True)
    values = list(workbook[sheet].values)
    records = [dict(zip(values[0], row, strict=True)) for row in values[1:]]
    return workbook, {row["parameter"]: row for row in records}


@pytest.mark.parametrize(
    ("mode", "expected_dispatch"),
    [
        ("DA MILP Replay", "DA-only MILP multi-cycle"),
        ("DA + IDA1 Replay", "Two-stage DA+IDA1 MILP multi-cycle"),
    ],
)
def test_multiday_saved_workbook_has_its_solver_and_wear_basis(
    mode: str, expected_dispatch: str,
) -> None:
    index = pd.date_range("2025-06-01", periods=48, freq="h", tz="UTC")
    prices = [10.0] * 6 + [100.0] * 6 + [20.0] * 6 + [90.0] * 6
    da = pd.DataFrame({"price_eur_mwh": prices * 2}, index=index)
    ida = pd.DataFrame({"intraday_price_eur_mwh": [p + 5 for p in prices] * 2}, index=index)
    assumptions = _global_assumptions()
    original = assumptions.copy(deep=True)

    bundle = cockpit._compute_multi_day_bundle(
        fingerprint="frozen-input", primary_df=da, intraday_df=ida,
        batch_dates=[pd.Timestamp("2025-06-01").date(), pd.Timestamp("2025-06-02").date()],
        mode=mode, zone_tz="UTC", carry_soc=True, power_mw=1.0,
        duration_hours=2, efficiency=0.88, capture_rate=1.0,
        capex_eur_kwh=150.0, assumptions=assumptions,
    )
    batch = bundle["batch"]
    assert len(batch) == 2
    assert batch["total_revenue_eur"].sum() > 0
    assert batch["degradation_cost_eur"].sum() > 0
    workbook, rows = _workbook_rows(cockpit_tables_to_excel(
        {"Multi-day Replay": batch}, assumptions=bundle["export_assumptions"],
    ), "Assumptions")

    assert rows["Dispatch model"]["value"] == expected_dispatch
    assert rows["Dispatch model"]["source"] == "Multi-day replay"
    assert rows["CapEx"]["value"] == "150"
    assert "degradation" in rows["CapEx"]["affects"].lower()
    assert "dispatch" in rows["CapEx"]["affects"].lower()
    assert rows["Cockpit capture haircut"]["value"] == "100%"
    assert "Capture rate (DA slippage)" not in rows
    assert len([row for row in workbook["Assumptions"].values if row[0] == "Dispatch model"]) == 1
    result_sheet = workbook["Multi-day Replay"]
    headers = [cell.value for cell in result_sheet[1]]
    for column in ("total_revenue_eur", "degradation_cost_eur", "daily_fce"):
        cell = result_sheet.cell(2, headers.index(column) + 1)
        assert cell.data_type == "n"
        # The existing workbook writer stores two-decimal numeric values.
        assert cell.value == pytest.approx(round(float(batch.iloc[0][column]), 2))
    pd.testing.assert_frame_equal(assumptions, original)
    assert assumptions.set_index("parameter").loc["Dispatch model", "value"] == "Greedy single-cycle"
    assert assumptions.set_index("parameter").loc["CapEx", "affects"] == "Payback period only (0 = skipped)"


def test_forecast_saved_workbook_describes_sequential_milp() -> None:
    assumptions = _global_assumptions()
    original = assumptions.copy(deep=True)
    export_assumptions = cockpit._forecast_policy_export_assumptions(
        assumptions, reserve_total=None, reserve_product=None,
        reserve_price=None, triple={"triple_total": None, "realistic_total": None},
        stochastic={"summary": None},
    )
    workbook, rows = _workbook_rows(cockpit_tables_to_excel(
        {"Strategy comparison": pd.DataFrame({"window_revenue_eur": [123.45]})},
        assumptions=export_assumptions,
    ), "Assumptions")

    assert rows["Dispatch model"]["value"] == "Sequential DA+IDA1 MILP multi-cycle"
    assert rows["Dispatch model"]["source"] == "Forecast-policy panel"
    assert "optional" in rows["Dispatch model"]["affects"].lower()
    assert rows["Capture haircut"]["value"] == "not applied"
    assert workbook["Strategy comparison"]["A2"].data_type == "n"
    assert workbook["Strategy comparison"]["A2"].value == 123.45
    pd.testing.assert_frame_equal(assumptions, original)


def test_missing_sidebar_dispatch_row_gets_one_panel_basis() -> None:
    assumptions = _global_assumptions()
    assumptions = assumptions.loc[assumptions["parameter"] != "Dispatch model"].copy()
    export = cockpit._forecast_policy_export_assumptions(
        assumptions, reserve_total=None, reserve_product=None,
        reserve_price=None, triple={"triple_total": None, "realistic_total": None},
        stochastic={"summary": None},
    )
    rows = export.loc[export["parameter"] == "Dispatch model"]
    assert rows["value"].tolist() == ["Sequential DA+IDA1 MILP multi-cycle"]
    assert "Dispatch model" not in set(assumptions["parameter"])


def test_absent_global_assumptions_stay_absent() -> None:
    assert cockpit._forecast_policy_export_assumptions(
        None, reserve_total=None, reserve_product=None,
        reserve_price=None, triple={"triple_total": None, "realistic_total": None},
        stochastic={"summary": None},
    ) is None


def test_global_and_cockpit_adapters_share_dispatch_and_capex_labels(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src import assumptions as assumption_module

    monkeypatch.setattr(assumption_module, "DISPATCH_PARAM_LABEL", "Dispatch solver", raising=False)
    monkeypatch.setattr(assumption_module, "CAPEX_PARAM_LABEL", "Capital cost", raising=False)
    monkeypatch.setattr(cockpit, "DISPATCH_PARAM_LABEL", "Dispatch solver", raising=False)
    monkeypatch.setattr(cockpit, "CAPEX_PARAM_LABEL", "Capital cost", raising=False)
    global_table = _global_assumptions()
    export_table = cockpit._frontier_basis_export_assumptions(
        global_table, capex_eur_kwh=150.0,
    )
    assert global_table["parameter"].tolist().count("Dispatch solver") == 1
    assert global_table["parameter"].tolist().count("Capital cost") == 1
    assert export_table["parameter"].tolist().count("Dispatch solver") == 1
    assert export_table["parameter"].tolist().count("Capital cost") == 1
    assert "Dispatch model" not in set(export_table["parameter"])
    assert "CapEx" not in set(export_table["parameter"])
