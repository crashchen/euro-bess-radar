"""Run-state and Excel-provenance regressions for frontier and floor panels."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import src.pages.simulation_cockpit as cockpit
from src.assumptions import CAPTURE_PARAM_LABEL, build_assumptions_table


def _prices() -> pd.DataFrame:
    day = [10.0] * 6 + [100.0] * 3 + [10.0] * 6 + [100.0] * 3 + [50.0] * 6
    index = pd.date_range(
        "2026-03-02", periods=96, freq="h", tz="UTC", name="timestamp",
    )
    return pd.DataFrame({"price_eur_mwh": day * 4}, index=index)


def _panel_app() -> None:
    import streamlit as st

    from src.pages.simulation_cockpit import (
        _render_contracted_floor_section,
        _render_cycle_frontier_section,
    )

    theme = st.checkbox("Light theme", key="followup_light_theme")
    context = _render_cycle_frontier_section(
        primary_zone="DE_LU",
        primary_df=st.session_state["followup_prices"],
        dates=st.session_state["followup_dates"],
        zone_tz="UTC",
        power_mw=1.0,
        duration_hours=1.0,
        efficiency=0.9,
        capex_eur_kwh=150.0,
        chart_template="plotly_white" if theme else "plotly_dark",
        assumptions=st.session_state["followup_assumptions"],
    )
    st.session_state["followup_frontier_context"] = context
    _render_contracted_floor_section(
        frontier_context=context,
        chart_template="plotly_white" if theme else "plotly_dark",
        assumptions=st.session_state["followup_assumptions"],
    )


def _elements(node, kind: str) -> list:
    found = []
    children = getattr(node, "children", None)
    if isinstance(children, dict):
        for child in children.values():
            if getattr(child, "type", "") == kind:
                found.append(child)
            found.extend(_elements(child, kind))
    return found


@pytest.fixture()
def app() -> AppTest:
    panel = AppTest.from_function(_panel_app)
    panel.session_state["followup_prices"] = _prices()
    panel.session_state["followup_dates"] = [
        pd.Timestamp(day).date()
        for day in ("2026-03-02", "2026-03-03", "2026-03-05")
    ]
    assumptions = build_assumptions_table(
        power_mw=1.0, duration_hours=1.0, efficiency=0.9,
        capture_rate=0.7, capex_eur_kwh=150.0, use_lp_dispatch=False,
    )
    assumptions.loc[len(assumptions)] = {
        "parameter": "Harness provenance", "value": "at-run", "unit": "",
        "source": "Test input", "affects": "Export audit",
    }
    panel.session_state["followup_assumptions"] = assumptions
    panel.session_state["cycle_frontier_caps"] = ["1", "uncapped"]
    panel.run(timeout=30)
    assert not panel.exception
    return panel


def _run_frontier(app: AppTest) -> None:
    app.button(key="cycle_frontier_run").click().run(timeout=60)
    assert not app.exception


def _run_floor(app: AppTest) -> None:
    app.button(key="contracted_floor_run").click().run(timeout=30)
    assert not app.exception


def _download_labels(app: AppTest) -> list[str]:
    return [item.proto.label for item in _elements(app.main, "download_button")]


def _workbook_rows(data: bytes, sheet: str) -> list[tuple]:
    book = load_workbook(BytesIO(data), data_only=True)
    assert sheet in book.sheetnames
    return list(book[sheet].values)


def _assumption_rows_by_parameter(rows: list[tuple]) -> dict[str, dict]:
    labels = rows[0]
    records = (dict(zip(labels, row, strict=True)) for row in rows[1:])
    return {record["parameter"]: record for record in records}


def test_floor_excel_uses_panel_basis_and_run_time_snapshot(
    app: AppTest, monkeypatch: pytest.MonkeyPatch,
) -> None:
    exported: list[bytes] = []
    frontier_exports: list[bytes] = []
    floor_solves: list[dict] = []
    real_export = cockpit.cockpit_tables_to_excel
    real_floor = cockpit.compute_decaying_contracted_floor_overlay

    def record_export(tables, *, assumptions=None):
        data = real_export(tables, assumptions=assumptions)
        if "Contracted floor" in tables:
            exported.append(data)
        if "Cycle-cap frontier" in tables:
            frontier_exports.append(data)
        return data

    def record_floor(**kwargs):
        floor_solves.append(dict(kwargs))
        return real_floor(**kwargs)

    monkeypatch.setattr(cockpit, "cockpit_tables_to_excel", record_export)
    monkeypatch.setattr(cockpit, "compute_decaying_contracted_floor_overlay", record_floor)
    _run_frontier(app)
    frontier_assumptions = _assumption_rows_by_parameter(
        _workbook_rows(frontier_exports[-1], "Assumptions")
    )
    assert frontier_assumptions["Dispatch model"]["value"] == (
        "DA-only MILP multi-cycle"
    )
    assert frontier_assumptions["CapEx"]["value"] == "150"
    assert frontier_assumptions["CapEx"]["source"] == "Frontier (sidebar CapEx)"
    assert "linear wear" in frontier_assumptions["CapEx"]["affects"].lower()
    _run_floor(app)
    assert len(floor_solves) == 1
    assert len(exported) == 1
    initial = _workbook_rows(exported[-1], "Assumptions")
    by_parameter = _assumption_rows_by_parameter(initial)
    assert CAPTURE_PARAM_LABEL not in by_parameter
    assert by_parameter["Cockpit capture haircut"]["value"] == "Not applied"
    assert "DA-only frontier" in by_parameter["Cockpit capture haircut"]["affects"]
    assert by_parameter["Dispatch model"]["value"] == "DA-only MILP multi-cycle"
    assert "linear wear" in by_parameter["CapEx"]["affects"].lower()
    assert by_parameter["CapEx"]["value"] == "150"
    assert by_parameter["CapEx"]["source"] == "Frontier (sidebar CapEx)"
    assert by_parameter["Harness provenance"]["value"] == "at-run"
    assert app.session_state["followup_assumptions"].set_index("parameter").loc[
        "Dispatch model", "value"
    ] == "Greedy single-cycle"
    floor_rows = _workbook_rows(exported[-1], "Contracted floor")
    merchant_column = floor_rows[0].index("merchant_net_eur_per_mw_yr")
    assert isinstance(floor_rows[1][merchant_column], (int, float))

    # Mutating the caller-owned global frame must not rewrite a saved result's
    # Assumptions sheet or silently rerun the floor calculation.
    global_table = app.session_state["followup_assumptions"]
    global_table.loc[
        global_table["parameter"] == "Harness provenance", "value"
    ] = "after-run"
    app.checkbox(key="followup_light_theme").check().run(timeout=30)
    assert not app.exception
    assert len(floor_solves) == 1
    assert len(exported) == 2
    assert _workbook_rows(exported[-1], "Assumptions") == initial
    assert _workbook_rows(exported[-1], "Contracted floor") == floor_rows

    # A new explicit Run takes a fresh snapshot without changing the solver's
    # inputs or the workbook's numeric result types.
    _run_floor(app)
    assert len(floor_solves) == 2
    updated = _workbook_rows(exported[-1], "Assumptions")
    updated_by_parameter = _assumption_rows_by_parameter(updated)
    assert updated_by_parameter["Harness provenance"]["value"] == "after-run"
    assert _workbook_rows(exported[-1], "Contracted floor") == floor_rows


def test_failed_frontier_retry_clears_previous_success(
    app: AppTest, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _run_frontier(app)
    assert any("frontier" in label.lower() for label in _download_labels(app))
    assert app.session_state["followup_frontier_context"] is not None
    attempts = []

    def fail(*args, **kwargs):
        attempts.append(1)
        raise ValueError("injected retry failure")

    monkeypatch.setattr(cockpit, "compute_cycle_cap_frontier", fail)
    app.button(key="cycle_frontier_run").click().run(timeout=30)
    assert not app.exception
    assert len(attempts) == 1
    assert any("Frontier sweep failed" in item.value for item in app.error)
    assert app.session_state["followup_frontier_context"] is None
    assert not _download_labels(app)

    app.run(timeout=30)
    assert not app.exception
    assert len(attempts) == 1
    assert app.session_state["followup_frontier_context"] is None
    assert not app.dataframe
    assert not _download_labels(app)


def test_pre_snapshot_floor_session_requires_a_fresh_run(app: AppTest) -> None:
    _run_frontier(app)
    _run_floor(app)
    cached = app.session_state["contracted_floor_result"]
    app.session_state["contracted_floor_result"] = {
        "fingerprint": cached["fingerprint"][1:],
        "result": cached["result"],
    }
    app.run(timeout=30)
    assert not app.exception
    assert any("Contract or frontier inputs changed" in item.value for item in app.info)
    assert "Annual merchant net" not in [metric.label for metric in app.metric]
    assert len(_download_labels(app)) == 1


def test_pre_provenance_frontier_session_requires_a_fresh_run(app: AppTest) -> None:
    _run_frontier(app)
    cached = app.session_state["cycle_frontier_result"]
    app.session_state["cycle_frontier_result"] = {
        **cached,
        "fingerprint": (
            "cockpit-cycle-frontier/v2-content-snapshot",
            *cached["fingerprint"][1:],
        ),
    }
    app.run(timeout=30)
    assert not app.exception
    assert app.session_state["followup_frontier_context"] is None
    assert any("Inputs changed since the last sweep" in item.value for item in app.info)
    assert not _download_labels(app)


def test_failed_floor_retry_clears_previous_success(
    app: AppTest, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _run_frontier(app)
    _run_floor(app)
    assert len(_download_labels(app)) == 2
    attempts = []

    def fail(**kwargs):
        attempts.append(1)
        raise ValueError("injected retry failure")

    monkeypatch.setattr(cockpit, "compute_decaying_contracted_floor_overlay", fail)
    app.button(key="contracted_floor_run").click().run(timeout=30)
    assert not app.exception
    assert len(attempts) == 1
    assert any("Contracted-floor calculation failed" in item.value for item in app.error)
    assert len(_download_labels(app)) == 1

    app.run(timeout=30)
    assert not app.exception
    assert len(attempts) == 1
    assert "Annual merchant net" not in [metric.label for metric in app.metric]
    assert len(_download_labels(app)) == 1
